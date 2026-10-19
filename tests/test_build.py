import datetime as dt
from contextlib import closing

import openpyxl
import pytest

from openpyxl.utils import column_index_from_string

from papaki_dre.build import SUM_COLUMNS, SheetReport, build_workbook, formula_coverage
from papaki_dre.types import ContractError, Formula, Sanitized
from papaki_dre.xlsx_inject import WorkbookSession, changed_parts
from xlsx_factory import make_template

D = dt.datetime(2026, 3, 17)
COM_HEADER = ["Modelo", "Núm. JOB", "Data", *[f"col{j}" for j in range(4, 15)]]   # 14 colunas (A:N)


def _san(report, header, rows, title=None):
    return Sanitized(report=report, header=header, rows=rows, db_rows=rows,
                     source_rows=list(range(len(rows))), raw_count=len(rows), title=title)


def _com_row(i, with_formulas=True):
    """Linha de comissões como o sanitizador entrega: A:K valores, L:N fórmulas de origem
    já reendereçadas para a linha de destino (2 + i)."""
    r = 2 + i
    head = ["M", f"#{i}", D, *([None] * 6), 100.0, 10.0]                             # A:K
    tail = ([Formula(f"=J{r}-K{r}"), Formula(f"=L{r}/J{r}"), Formula(f"=K{r}*2")]   # L:N
            if with_formulas else [90.0, 0.9, 20.0])
    return head + tail


def _results(n_rra=2, n_cp=3, n_cr=2, n_com=2):
    return {
        "rra": _san("rra", ["Núm. JOB", "Data", "VALOR"],
                    [[f"75{i}", D, 100.0 + i] for i in range(n_rra)], title="RENDIMENTOS DOS JOBS"),
        "cp": _san("cp", ["Cód.", "Valor", "Vcto"], [[f"x-{i}", 10.5 + i, D] for i in range(n_cp)]),
        "cr": _san("cr", ["Cliente", "Emissão", "Vl"], [[f"c{i}", D, 5.0] for i in range(n_cr)]),
        "comissoes": _san("comissoes", COM_HEADER, [_com_row(i) for i in range(n_com)]),
    }


def _wb(path, **kw):
    """openpyxl.load_workbook como gerenciador de contexto (fecha o arquivo no fim)."""
    return closing(openpyxl.load_workbook(path, **kw))


def _build(tmp_path, results=None, **tpl):
    tpl_path = make_template(tmp_path / "t.xlsx", **tpl)
    out = tmp_path / "o.xlsx"
    report = build_workbook(tpl_path, out, results or _results())
    return tpl_path, out, report


def test_writes_headers_title_and_data_in_the_right_places(tmp_path):
    _, out, _ = _build(tmp_path)
    with _wb(out) as wb:
        base = wb["BASE"]
        assert base["A1"].value == "RENDIMENTOS DOS JOBS"
        assert [c.value for c in base[3][:3]] == ["Núm. JOB", "Data", "VALOR"]
        assert base["A5"].value == "750" and base["B5"].value == D and base["C6"].value == 101.0
        desp = wb["DESPESAS"]
        assert [c.value for c in desp[2][:3]] == ["Cód.", "Valor", "Vcto"]
        assert desp["A3"].value == "x-0" and desp["C5"].value == D
        assert wb["RECEITA"]["A1"].value == "Cliente" and wb["RECEITA"]["A3"].value == "c1"
        com = wb["REAL X ORÇADO COMISSÕES"]
        assert com["A1"].value == "Modelo" and com["B2"].value == "#0" and com["C3"].value == D


def test_base_total_row_is_two_rows_below_the_last_job(tmp_path):
    _, out, report = _build(tmp_path)
    with _wb(out) as wb:
        base = wb["BASE"]
        # 2 jobs: linhas 5 e 6 -> total na 8, somando da linha 4 até a 7
        assert base["A8"].value == "TOTAL"
        assert base["J8"].value == "=SUM(J4:J7)"
        assert base["AU8"].value == "=SUM(AU4:AU7)"
    assert len(SUM_COLUMNS) == 35 and SUM_COLUMNS[0] == "J" and SUM_COLUMNS[-1] == "AU"
    assert report.sheets["BASE"].total_row == 8
    assert report.sheets["BASE"].missing_formulas == []        # TOTAL completo na cobertura


def test_coverage_flags_a_missing_base_total(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx")                    # BASE do modelo: sem TOTAL
    missing = formula_coverage(tpl, "rra", SheetReport("BASE", 5, 6, 2, total_row=8))
    assert missing[0] == "A8"                                   # rótulo ausente
    assert missing[1:] == [f"{c}8" for c in SUM_COLUMNS]        # e as 35 somas


def test_coverage_flags_a_base_total_with_the_wrong_range(tmp_path):
    with WorkbookSession(make_template(tmp_path / "t.xlsx")) as s:
        ed = s.editor("BASE")
        ed.write(8, 1, "TOTAL")
        for letter in SUM_COLUMNS:
            end = 6 if letter == "J" else 7                     # J soma até a linha errada
            ed.write(8, column_index_from_string(letter), Formula(f"=SUM({letter}4:{letter}{end})"))
        out = s.save(tmp_path / "o.xlsx")
    assert formula_coverage(out, "rra", SheetReport("BASE", 5, 6, 2, total_row=8)) == ["J8"]


def test_formulas_of_the_model_are_extended_beyond_the_template(tmp_path):
    # o modelo tem fórmulas até a linha 7; carregando 8 linhas (3..10) precisa estender 8..10
    _, out, report = _build(tmp_path, _results(n_cp=8), desp_formula_rows=5)
    with WorkbookSession(out) as sess:
        ed = sess.editor("DESPESAS")
        assert ed.formula_at(10, 42) == '=IF(A10="x","a","b")'
        assert ed.formula_at(10, 43) == '=LEN(AP10)&"!"'
    assert report.sheets["DESPESAS"].missing_formulas == []
    assert report.sheets["DESPESAS"].orphan_formula_rows == 0


def test_orphan_formulas_below_loaded_rows_are_kept_and_counted(tmp_path):
    _, out, report = _build(tmp_path, _results(n_cp=2), desp_formula_rows=5)   # carrega 3..4; modelo vai até 7
    assert report.sheets["DESPESAS"].orphan_formula_rows == 3
    with WorkbookSession(out) as sess:
        assert sess.editor("DESPESAS").formula_at(7, 42) is not None


def test_receita_formula_column_is_extended(tmp_path):
    _, out, report = _build(tmp_path, _results(n_cr=6), rec_formula_rows=3)
    with WorkbookSession(out) as sess:
        assert sess.editor("RECEITA").formula_at(7, 29) == '=IF(A7="","",A7)'
    assert report.sheets["RECEITA"].missing_formulas == []


def test_comissoes_use_september_formulas_when_the_model_has_none(tmp_path):
    _, out, report = _build(tmp_path)
    with _wb(out) as wb:
        ws = wb["REAL X ORÇADO COMISSÕES"]
        assert ws["O3"].value == "=DATE(YEAR(C3),MONTH(C3),1)"
        assert ws["P3"].value.startswith("=IF(COUNTIFS($B$2:$B3,$B3,$H$2:$H3,$H3,$G$2:$G3,$G3,$C$2:$C3,$C3)=1,$J3,0)")
        assert ws["Q2"].value.startswith('=IF($H2="Diretor",$P2/')
    assert report.sheets["REAL X ORÇADO COMISSÕES"].missing_formulas == []


def test_comissoes_source_formulas_in_L_to_N_are_written_and_covered(tmp_path):
    _, out, report = _build(tmp_path)
    with _wb(out) as wb:
        ws = wb["REAL X ORÇADO COMISSÕES"]
        assert ws["L3"].value == "=J3-K3" and ws["N2"].value == "=K2*2"
    assert report.sheets["REAL X ORÇADO COMISSÕES"].missing_formulas == []


def test_comissoes_row_without_formulas_in_L_to_N_is_a_coverage_gap(tmp_path):
    res = _results()
    res["comissoes"].rows[1] = _com_row(1, with_formulas=False)     # linha 3 chegou só com valores
    _, _, report = _build(tmp_path, res)
    # build_workbook não levanta: a lacuna vai para missing_formulas, e quem consome o relatório
    # (a etapa de conferência) é que deve bloquear a publicação enquanto a lista não estiver vazia
    assert report.sheets["REAL X ORÇADO COMISSÕES"].missing_formulas == ["L3", "M3", "N3"]


def test_missing_model_formulas_block_the_publication(tmp_path):
    # classificação ausente: o modelo não tem AP:AQ para estender. O ContractError interrompe a carga
    # antes de gravar qualquer arquivo; a importação deve ser tratada como bloqueada, com o motivo.
    tpl = make_template(tmp_path / "t.xlsx", desp_formula_rows=0)
    with pytest.raises(ContractError, match="classificação ausente"):
        build_workbook(tpl, tmp_path / "o.xlsx", _results())
    assert not (tmp_path / "o.xlsx").exists()


def test_dirty_template_is_rejected_and_nothing_is_written(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx")
    with WorkbookSession(tpl) as s:
        s.editor("DESPESAS").write(3, 1, "resto de outra rodada")
        dirty = s.save(tmp_path / "dirty.xlsx")
    with pytest.raises(ContractError, match="A3"):
        build_workbook(dirty, tmp_path / "o.xlsx", _results())
    assert not (tmp_path / "o.xlsx").exists()


def test_only_the_four_base_sheets_change(tmp_path):
    tpl, out, report = _build(tmp_path)
    changed = set(report.parts["changed"])
    assert changed == {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml",
                       "xl/worksheets/sheet3.xml", "xl/worksheets/sheet4.xml",
                       "xl/workbook.xml", "[Content_Types].xml", "xl/_rels/workbook.xml.rels"}
    assert "xl/worksheets/sheet5.xml" not in changed          # aba "OUTRA" intacta
    assert report.parts["removed"] == ["xl/calcChain.xml"]


def test_date_in_unformatted_cells_is_reported(tmp_path):
    _, _, report = _build(tmp_path)
    # depende da fábrica: a RECEITA sintética não tem estilo nenhum nas colunas de dados, então as 2
    # datas da coluna B caem sem formato de data (nas outras abas a coluna das datas tem estilo de data)
    assert report.sheets["RECEITA"].date_style_warnings == {"B": 2}


def test_zero_rows_still_builds(tmp_path):
    res = _results(n_rra=0, n_cp=0, n_cr=0, n_com=0)
    _, out, report = _build(tmp_path, res)
    assert report.sheets["DESPESAS"].rows == 0
    with _wb(out) as wb:
        assert wb["BASE"]["A6"].value == "TOTAL"                       # sem jobs: 4 + 2


def test_typed_constants_inside_loaded_rows_are_replaced_by_the_model_formula(tmp_path):
    # modelo com 'Desconsiderar' digitado à mão em AP/AQ (linhas 4 e 6) e em RECEITA!AC (linha 3)
    _, out, report = _build(tmp_path, _results(n_cp=5, n_cr=3), desp_formula_rows=5,
                            rec_formula_rows=3, desp_constant_rows=(4, 6), rec_constant_rows=(3,))
    with WorkbookSession(out) as sess:
        desp, rec = sess.editor("DESPESAS"), sess.editor("RECEITA")
        for r in (4, 6):
            assert desp.formula_at(r, 42) == f'=IF(A{r}="x","a","b")'
            assert desp.formula_at(r, 43) == f'=LEN(AP{r})&"!"'
        assert rec.formula_at(3, 29) == '=IF(A3="","",A3)'
        for r in (3, 5, 7):                                      # linhas que já tinham fórmula: intactas
            assert desp.formula_at(r, 42) == f'=IF(A{r}="x","a","b")'
        assert rec.formula_at(2, 29) == '=IF(A2="","",A2)' and rec.formula_at(4, 29) == '=IF(A4="","",A4)'
    assert report.sheets["DESPESAS"].missing_formulas == []
    assert report.sheets["RECEITA"].missing_formulas == []
    assert report.sheets["DESPESAS"].overwritten_constants == 4   # AP e AQ em 2 linhas
    assert report.sheets["RECEITA"].overwritten_constants == 1
    assert report.sheets["BASE"].overwritten_constants == 0


def test_no_constants_means_nothing_overwritten(tmp_path):
    _, _, report = _build(tmp_path, _results(n_cp=8))             # estende 8..10: células vazias não contam
    assert all(sr.overwritten_constants == 0 for sr in report.sheets.values())


def test_replaced_constants_are_formulas_for_openpyxl_too(tmp_path):
    _, out, _ = _build(tmp_path, _results(n_cp=5), desp_constant_rows=(4,))
    with _wb(out) as wb:
        ws = wb["DESPESAS"]
        assert ws["AP4"].value == '=IF(A4="x","a","b")' and ws["AQ4"].value == '=LEN(AP4)&"!"'


def test_coverage_reuses_one_session_for_all_sheets(tmp_path, monkeypatch):
    import papaki_dre.build as build
    opened = []

    class Counting(WorkbookSession):
        def __init__(self, path):
            opened.append(path)
            super().__init__(path)

    monkeypatch.setattr(build, "WorkbookSession", Counting)
    tpl, out, _ = _build(tmp_path)
    assert opened == [tpl, out]                                  # modelo + UMA leitura da saída


def test_formula_coverage_with_and_without_an_open_session_agree(tmp_path):
    _, out, report = _build(tmp_path, _results(n_cp=8))
    sr = report.sheets["DESPESAS"]
    with WorkbookSession(out) as s:
        assert formula_coverage(out, "cp", sr, session=s) == formula_coverage(out, "cp", sr) == []


def test_summary_row_below_the_data_is_not_used_as_the_extension_base(tmp_path):
    # fórmulas de dados em 3..7 e uma linha de resumo na 9 (COUNTA em AP, $A$1 em AQ): carregando
    # 6 linhas (3..8), a 8 recebe a fórmula da família dos dados e o resumo não conta como órfã
    _, out, report = _build(tmp_path, _results(n_cp=6), desp_formula_rows=5, desp_footer_row=9)
    with WorkbookSession(out) as sess:
        ed = sess.editor("DESPESAS")
        assert ed.formula_at(8, 42) == '=IF(A8="x","a","b")'
        assert ed.formula_at(8, 43) == '=LEN(AP8)&"!"'
        assert ed.formula_at(9, 42) == "=COUNTA(AP3:AP8)"          # o resumo fica intacto
    sr = report.sheets["DESPESAS"]
    assert sr.missing_formulas == [] and sr.orphan_formula_rows == 0


def test_only_a_summary_formula_is_still_a_missing_classification(tmp_path):
    # sem fórmula de dados: o COUNTA do resumo não pode virar base de extensão de AP
    tpl = make_template(tmp_path / "t.xlsx", desp_formula_rows=0, desp_footer_row=6)
    with pytest.raises(ContractError, match="classificação ausente"):
        build_workbook(tpl, tmp_path / "o.xlsx", _results())


def test_base_total_row_copies_the_style_of_the_last_job_row(tmp_path):
    from xlsx_factory import replace_part, sheet_xml
    # BASE com estilo só na linha 5 (a do único job): A5 data (3), J5/AU5 dinheiro (2); linha 7 sem células
    rows = '<row r="5"><c r="A5" s="3"/><c r="J5" s="2"/><c r="AU5" s="2"/></row>'
    tpl = replace_part(make_template(tmp_path / "t.xlsx"), tmp_path / "b.xlsx",
                       "xl/worksheets/sheet1.xml", sheet_xml(rows))
    build_workbook(tpl, tmp_path / "o.xlsx", _results(n_rra=1))
    with WorkbookSession(tmp_path / "o.xlsx") as s:
        ed = s.editor("BASE")
        assert ed.text_at(7, 1) == "TOTAL" and ed.style_of(7, 1) == "3"
        assert ed.style_of(7, 10) == "2" and ed.style_of(7, 47) == "2"
        assert ed.formula_at(7, 10) == "=SUM(J4:J6)"


def test_base_sum_start_row_is_a_named_constant():
    from papaki_dre import build
    assert build.BASE_SUM_START_ROW == 4


def test_empty_or_unresolvable_formulas_are_coverage_gaps(tmp_path):
    from xlsx_factory import replace_part, sheet_xml
    rows = ('<row r="3"><c r="AP3"><f t="shared" ref="AP3:AP5" si="0">A3</f></c><c r="AQ3"><f>1</f></c></row>'
            '<row r="4"><c r="AP4"><f/></c><c r="AQ4"><f t="shared" si="9"/></c></row>'   # vazia | filha sem mestre
            '<row r="5"><c r="AP5"><f t="shared" si="0"/></c><c r="AQ5"><f> </f></c></row>')  # filha ok | só espaço
    tpl = make_template(tmp_path / "t.xlsx")
    path = replace_part(tpl, tmp_path / "b.xlsx", "xl/worksheets/sheet2.xml", sheet_xml(rows))
    assert formula_coverage(path, "cp", SheetReport("DESPESAS", 3, 5, 3)) == ["AP4", "AQ4", "AQ5"]


def _filter_state(out):
    import re
    import zipfile
    with zipfile.ZipFile(out) as z:
        sheet = z.read("xl/worksheets/sheet2.xml").decode()
        names = re.search(r"<definedNames>.*</definedNames>", z.read("xl/workbook.xml").decode()).group(0)
    return re.search(r'<autoFilter ref="([^"]+)"', sheet).group(1), names


OTHER_NAMES = ('<definedName name="_xlnm._FilterDatabase" localSheetId="4" hidden="1">OUTRA!$A$1:$B$1'
               '</definedName><definedName name="Taxa">OUTRA!$A$1</definedName>')


def test_autofilter_is_extended_to_the_last_loaded_row(tmp_path):
    # filtro do modelo em A2:AQ7; carregando 8 linhas (3..10) o filtro e o nome oculto vão até a 10
    _, out, report = _build(tmp_path, _results(n_cp=8), desp_autofilter="A2:AQ7")
    ref, names = _filter_state(out)
    assert ref == "A2:AQ10"
    assert ('<definedName name="_xlnm._FilterDatabase" localSheetId="1" hidden="1">'
            'DESPESAS!$A$2:$AQ$10</definedName>') in names
    assert names.endswith(OTHER_NAMES + "</definedNames>")         # outras abas e nomes: intactos
    with _wb(out) as wb:
        assert wb["DESPESAS"].auto_filter.ref == "A2:AQ10"
    assert report.sheets["DESPESAS"].autofilter_extended is True
    assert not any(sr.autofilter_extended for k, sr in report.sheets.items() if k != "DESPESAS")


def test_autofilter_is_never_shrunk(tmp_path):
    _, out, report = _build(tmp_path, _results(n_cp=2), desp_autofilter="A2:AQ7")   # carrega só 3..4
    ref, names = _filter_state(out)
    assert ref == "A2:AQ7"
    assert "DESPESAS!$A$2:$AQ$7</definedName>" in names
    assert report.sheets["DESPESAS"].autofilter_extended is False


def test_filter_name_already_larger_than_the_new_filter_is_not_shrunk():
    from papaki_dre.xlsx_inject import _sync_filter_names
    wb = (b'<workbook><definedNames><definedName name="_xlnm._FilterDatabase" localSheetId="1" hidden="1">'
          b"'DESP X'!$A$2:$AR$900</definedName></definedNames></workbook>")
    assert _sync_filter_names(wb, {1: "A2:AR500"}) == wb
    grown = _sync_filter_names(wb, {1: "A2:AR1200"})
    assert b"'DESP X'!$A$2:$AR$1200</definedName>" in grown


def test_errors_booleans_and_untrimmed_text_pass_through_build_workbook(tmp_path):
    from papaki_dre.types import XlError
    res = _results(n_cp=3)
    res["cp"].rows[0] = [XlError("#N/A"), True, "  dois espaços  "]
    res["cp"].rows[1] = ["linha 1\nlinha 2", False, XlError("#DIV/0!")]
    _, out, report = _build(tmp_path, res)
    with _wb(out) as wb:
        ws = wb["DESPESAS"]
        assert ws["A3"].value == "#N/A" and ws["A3"].data_type == "e"
        assert ws["B3"].value is True and ws["B4"].value is False
        assert ws["C3"].value == "  dois espaços  "                  # espaços nas pontas preservados
        assert ws["A4"].value == "linha 1\nlinha 2"                  # quebra de linha preservada
        assert ws["C4"].value == "#DIV/0!" and ws["C4"].data_type == "e"
    assert report.sheets["DESPESAS"].missing_formulas == []


def test_sheet_report_new_fields_have_defaults_and_keep_positional_construction():
    sr = SheetReport("X", 1, 2, 2)
    assert sr.overwritten_constants == 0 and sr.autofilter_extended is False


def test_summary_row_inside_the_loaded_rows_is_replaced_and_counted(tmp_path):
    # fórmulas de dados em 3..7 e um resumo (COUNTA em AP, $A$1&"x" em AQ) na 9; carregando 8 linhas
    # (3..10) o resumo fica DENTRO da faixa: vira a fórmula de dados e conta como substituída
    from lxml import etree
    tpl = make_template(tmp_path / "t.xlsx", desp_footer_row=9)
    with WorkbookSession(tpl) as s:
        ed = s.editor("DESPESAS")
        before = {r: etree.tostring(ed._cell(r, 42, create=False)) for r in range(3, 8)}
    report = build_workbook(tpl, tmp_path / "o.xlsx", _results(n_cp=8))
    with WorkbookSession(tmp_path / "o.xlsx") as s:
        ed = s.editor("DESPESAS")
        assert ed.formula_at(9, 42) == '=IF(A9="x","a","b")'
        assert ed.formula_at(9, 43) == '=LEN(AP9)&"!"'
        c = ed._cell(9, 43, create=False)
        assert c.get("t") is None and len(c) == 1                    # sem o valor em cache do resumo
        for r in range(3, 8):                                        # mestre e filhas legítimas intactos
            assert etree.tostring(ed._cell(r, 42, create=False)) == before[r]
    sr = report.sheets["DESPESAS"]
    assert sr.overwritten_constants == 2 and sr.missing_formulas == [] and sr.orphan_formula_rows == 0


# ---- fórmula completa em toda a coluna (decisão do usuário de 13/10/2026) -----------------------

def _sheet_part(path, name="xl/worksheets/sheet2.xml"):
    import zipfile
    with zipfile.ZipFile(path) as z:
        return z.read(name)


def _cells_xml(path, col, rows, sheet="DESPESAS"):
    from lxml import etree
    with WorkbookSession(path) as s:
        ed = s.editor(sheet)
        return {r: (None if (c := ed._cell(r, col, create=False)) is None else etree.tostring(c))
                for r in rows}


def _full(r):
    return f'=IF(B{r}="t","Desconsiderar",IF(A{r}="x","a","b"))'


def test_a_linha_mestre_da_untitled_e_a_ap_completa_do_modelo():
    from papaki_dre.produtoras import carregar
    assert carregar("untitled").linha_mestre == {"cp": {"AP": 881}}


def test_build_com_a_linha_mestre_da_produtora_exige_a_linha_no_modelo(tmp_path):
    from papaki_dre.produtoras import carregar
    tpl = make_template(tmp_path / "t.xlsx")                    # sintético: AP só até a linha 7
    with pytest.raises(ContractError, match=r"DESPESAS!AP881.*MASTER_FORMULA_ROWS.*não tem fórmula"):
        build_workbook(tpl, tmp_path / "o.xlsx", _results(),
                       master_formula_rows=carregar("untitled").linha_mestre)
    assert not (tmp_path / "o.xlsx").exists()


def test_complete_formula_is_applied_to_every_row_of_the_configured_column(tmp_path):
    # AP: curta (si=0) em 3..4, COMPLETA (si=1) em 5..7, curta (si=2) em 8..10; constantes em AP/AQ nas
    # linhas 4 e 9. Carregando 6 linhas (3..8), AP inteira (3..10, incluindo as órfãs 9..10) fica com a
    # fórmula completa da linha 5; o grupo si=1 (já completo) fica byte a byte; AQ segue a regra de hoje
    tpl = make_template(tmp_path / "t.xlsx", desp_formula_rows=8, desp_ap_variant=(5, 7),
                        desp_constant_rows=(4, 9), desp_autofilter="A2:AQ10")
    ap_before = _cells_xml(tpl, 42, range(3, 11))
    aq_before = _cells_xml(tpl, 43, range(3, 11))
    out = tmp_path / "o.xlsx"
    report = build_workbook(tpl, out, _results(n_cp=6), master_formula_rows={"cp": {"AP": 5}})
    with WorkbookSession(out) as s:
        ed = s.editor("DESPESAS")
        assert [ed.formula_at(r, 42) for r in range(3, 11)] == [_full(r) for r in range(3, 11)]
        assert [ed.formula_at(r, 43) for r in (3, 4, 5, 8)] == [f'=LEN(AP{r})&"!"' for r in (3, 4, 5, 8)]
    ap_after, aq_after = _cells_xml(out, 42, range(3, 11)), _cells_xml(out, 43, range(3, 11))
    assert all(ap_after[r] == ap_before[r] for r in (5, 6, 7))          # já completas: intactas
    assert all(ap_after[r] != ap_before[r] for r in (3, 4, 8, 9, 10))
    assert all(aq_after[r] == aq_before[r] for r in (3, 5, 6, 7, 8, 9, 10))   # AQ: só a constante da 4
    xml = _sheet_part(out)
    from xlsx_factory import orphan_shared_children
    assert orphan_shared_children(xml) == []
    assert b'<autoFilter ref="A2:AQ10"/>' in xml
    from lxml import etree
    for row in etree.fromstring(xml).iter("{%s}row" % "http://schemas.openxmlformats.org/spreadsheetml/2006/main"):
        cols = [column_index_from_string("".join(ch for ch in c.get("r") if ch.isalpha())) for c in row]
        assert cols == sorted(set(cols))                             # células em ordem, sem repetição
    sr = report.sheets["DESPESAS"]
    assert sr.unified_formula_rows == 3                              # AP3 (mestre si=0), AP8 (mestre si=2), AP10
    assert sr.overwritten_constants == 3                             # AP4, AP9 e AQ4
    assert sr.orphan_formula_rows == 2 and sr.missing_formulas == []
    assert all(x.unified_formula_rows == 0 for k, x in report.sheets.items() if k != "DESPESAS")
    with _wb(out) as wb:
        assert wb["DESPESAS"]["AP10"].value == _full(10)


def test_configured_row_without_formula_is_a_contract_error(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx", desp_ap_variant=(4, 5))
    with pytest.raises(ContractError, match=r"DESPESAS!AP2: a linha configurada em MASTER_FORMULA_ROWS não tem fórmula"):
        build_workbook(tpl, tmp_path / "o.xlsx", _results(), master_formula_rows={"cp": {"AP": 2}})
    assert not (tmp_path / "o.xlsx").exists()


def test_configured_row_whose_formula_does_not_translate_to_first_is_a_contract_error(tmp_path):
    # AP9 = COUNTA(AP3:AP8): traduzida para a linha 3 cairia em AP-3
    tpl = make_template(tmp_path / "t.xlsx", desp_footer_row=9)
    with pytest.raises(ContractError, match=r"DESPESAS!AP9: a fórmula da linha configurada .* linha 3"):
        build_workbook(tpl, tmp_path / "o.xlsx", _results(), master_formula_rows={"cp": {"AP": 9}})
    assert not (tmp_path / "o.xlsx").exists()


def test_shared_group_not_fully_covered_by_the_unification_is_a_contract_error(tmp_path):
    # o mestre AP3 (curto) tem ref AP3:AR4: AR3/AR4 são filhas fora da coluna unificada (e AQ3/AQ4,
    # também no ref, não estão no lote de AP). Trocar o mestre deixaria AR3/AR4 sem mestre -> recusa
    # pelo ref inteiro, e nada é gravado
    from xlsx_factory import replace_part, sheet_xml
    rows = ('<row r="2"><c r="AP2" t="inlineStr"><is><t>De-Para</t></is></c></row>'
            '<row r="3"><c r="AP3"><f t="shared" ref="AP3:AR4" si="0">IF(A3="x","a","b")</f></c>'
            '<c r="AQ3"><f>LEN(AP3)</f></c><c r="AR3"><f t="shared" si="0"/></c></row>'
            '<row r="4"><c r="AP4"><f t="shared" si="0"/></c><c r="AQ4"><f>LEN(AP4)</f></c>'
            '<c r="AR4"><f t="shared" si="0"/></c></row>'
            '<row r="5"><c r="AP5"><f>IF(B5="t","Desconsiderar",IF(A5="x","a","b"))</f></c>'
            '<c r="AQ5"><f>LEN(AP5)</f></c></row>')
    tpl = replace_part(make_template(tmp_path / "t.xlsx"), tmp_path / "b.xlsx",
                       "xl/worksheets/sheet2.xml", sheet_xml(rows, "AP2:AR5"))
    with pytest.raises(ContractError, match=r"DESPESAS!AP3 é o mestre da fórmula compartilhada si=0 "
                                            r"\(ref AP3:AR4\).*de fora: AQ3, AR3, AQ4, AR4$"):
        build_workbook(tpl, tmp_path / "o.xlsx", _results(), master_formula_rows={"cp": {"AP": 5}})
    assert not (tmp_path / "o.xlsx").exists()


def test_unified_formula_rows_has_a_default_and_comes_last():
    import dataclasses
    assert SheetReport("X", 1, 2, 2).unified_formula_rows == 0
    assert [f.name for f in dataclasses.fields(SheetReport)][-1] == "unified_formula_rows"
