import datetime as dt
import sys
import zipfile
from contextlib import closing

import pytest

from papaki_dre.xlsx_inject import DateStyles, to_serial
from xlsx_factory import STYLES, make_template


def test_date_styles_recognise_builtin_and_custom_date_formats():
    ds = DateStyles(STYLES.encode("utf-8"))
    assert not ds.is_date(None)       # estilo 0: geral
    assert ds.is_date("1")            # numFmtId 14
    assert not ds.is_date("2")        # numFmtId 4 (dinheiro)
    assert ds.is_date("3")            # custom dd/mm/yyyy
    assert not ds.is_date("99")       # fora da tabela


def test_date_styles_without_styles_part():
    assert not DateStyles(None).is_date("1")


def test_to_serial_known_values():
    assert to_serial(dt.datetime(2026, 1, 1)) == 46023
    assert to_serial(dt.datetime(1900, 3, 1)) == 61
    assert to_serial(dt.datetime(2026, 1, 1, 12)) == pytest.approx(46023.5)
    assert to_serial(dt.date(2026, 1, 2)) == 46024


from lxml import etree

from papaki_dre.types import ContractError, Formula, XlError
from papaki_dre.xlsx_inject import NS, SheetEditor, WorkbookSession, changed_parts


@pytest.fixture
def editor(tmp_path):
    """editor(aba, **kw_do_modelo) -> (sessão, editor da aba); as sessões fecham no fim do teste."""
    opened = []

    def make(sheet="BASE", **kw):
        s = WorkbookSession(make_template(tmp_path / "t.xlsx", **kw))
        opened.append(s)
        return s, s.editor(sheet)

    yield make
    for s in opened:
        s.close()


def _wb(path, **kw):
    """openpyxl.load_workbook como gerenciador de contexto (fecha o arquivo no fim)."""
    import openpyxl
    return closing(openpyxl.load_workbook(path, **kw))


def _read_part(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name)


def test_write_all_value_types_and_read_back_with_openpyxl(tmp_path, editor):
    s, ed = editor()
    ed.write(5, 1, "texto com <&> e acentuação ç", date_styles=s.date_styles)
    ed.write(5, 2, dt.datetime(2026, 3, 17), date_styles=s.date_styles)
    ed.write(5, 3, 1234.5, date_styles=s.date_styles)
    ed.write(6, 1, 7, date_styles=s.date_styles)
    ed.write(6, 2, True)
    ed.write(6, 3, Formula("=A6*2"))
    ed.write(7, 1, XlError("#DIV/0!"))
    s.save(tmp_path / "o.xlsx")
    with _wb(tmp_path / "o.xlsx") as wb:
        ws = wb["BASE"]
        assert ws["A5"].value == "texto com <&> e acentuação ç"
        assert ws["B5"].value == dt.datetime(2026, 3, 17)       # estilo de data preservado
        assert ws["C5"].value == 1234.5
        assert ws["A6"].value == 7 and ws["B6"].value is True
        assert ws["C6"].value == "=A6*2"
        assert ws["A7"].value == "#DIV/0!" and ws["A7"].data_type == "e"


def test_none_writes_nothing_and_creates_no_row(tmp_path, editor):
    _, ed = editor()
    ed.write(500, 1, None)
    assert not ed.row_exists(500)


def test_existing_style_is_kept_and_new_cells_copy_style_from_reference(tmp_path, editor):
    s, ed = editor("DESPESAS")
    ed.write(3, 3, "x")                              # C3 já existe com s=3
    assert ed.style_of(3, 3) == "3"
    ed.write(20, 3, "y", style_from=(3, 3))           # linha nova copia o estilo de C3
    assert ed.style_of(20, 3) == "3"


def test_rows_and_cells_stay_in_order_when_inserting(tmp_path, editor):
    s, ed = editor("DESPESAS")
    ed.write(15, 5, "e")
    ed.write(15, 1, "a")
    ed.write(10, 2, "b")
    out = tmp_path / "o.xlsx"
    s.save(out)
    import re
    xml = _read_part(out, "xl/worksheets/sheet2.xml").decode()
    rows = [int(x) for x in re.findall(r'<row r="(\d+)"', xml)]
    assert rows == sorted(rows)
    assert re.findall(r'<c r="([A-Z]+)15"', xml) == ["A", "E"]


def test_refuses_to_overwrite_a_model_formula(tmp_path, editor):
    _, ed = editor("DESPESAS")
    with pytest.raises(ContractError, match="AP3"):
        ed.write(3, 42, "x")


def test_illegal_xml_characters_are_removed(tmp_path, editor):
    s, ed = editor()
    ed.write(5, 1, "a\x00b\x0bc")
    s.save(tmp_path / "o.xlsx")
    with _wb(tmp_path / "o.xlsx") as wb:
        assert wb["BASE"]["A5"].value == "abc"


def test_date_in_cell_without_date_format_is_counted(tmp_path, editor):
    s, ed = editor("RECEITA")
    ed.write(2, 2, dt.datetime(2026, 1, 5), date_styles=s.date_styles)
    assert ed.date_warnings == {"B": 1}


def test_nan_is_rejected(tmp_path, editor):
    _, ed = editor()
    with pytest.raises(ContractError):
        ed.write(5, 1, float("nan"))


def test_unsupported_type_is_rejected(tmp_path, editor):
    _, ed = editor()
    with pytest.raises(ContractError):
        ed.write(5, 1, object())


def test_nonempty_cells_detects_dirty_region(tmp_path, editor):
    _, ed = editor("DESPESAS")
    assert ed.nonempty_cells(range(3, 8), {1, 2, 3}) == []
    ed.write(4, 1, "sujeira")
    assert ed.nonempty_cells(range(3, 8), {1, 2, 3}) == ["A4"]


def test_formula_at_resolves_shared_formula_children(tmp_path, editor):
    _, ed = editor("DESPESAS")
    assert ed.formula_at(3, 42) == '=IF(A3="x","a","b")'     # mestre
    assert ed.formula_at(5, 42) == '=IF(A5="x","a","b")'     # filha: traduzida para a linha 5
    assert ed.formula_at(5, 43) == '=LEN(AP5)&"!"'            # explícita
    assert ed.formula_at(4, 1) is None


def test_text_at_reads_inline_strings_only(tmp_path, editor):
    _, ed = editor()
    ed.write(8, 1, "TOTAL")
    ed.write(9, 1, 5)
    assert ed.text_at(8, 1) == "TOTAL"
    assert ed.text_at(9, 1) is None          # número
    assert ed.text_at(500, 1) is None        # célula inexistente


def test_last_formula_row(tmp_path, editor):
    _, ed = editor("DESPESAS", desp_formula_rows=5)
    assert ed.last_formula_row(42) == 7
    assert ed.last_formula_row(43) == 7
    assert ed.last_formula_row(1) is None


def test_extension_writes_explicit_formulas_translated_from_last_row(tmp_path, editor):
    from papaki_dre.build import SheetReport, _extend_formulas
    s, ed = editor("DESPESAS", desp_formula_rows=5)          # fórmulas em 3..7
    report = SheetReport("DESPESAS", 3, 9, 7)
    _extend_formulas(ed, "cp", "DESPESAS", "AP", 3, 9, report)
    for r in (8, 9):
        f = ed._cell(r, 42, create=False).find("{%s}f" % NS)
        assert f.text == f'IF(A{r}="x","a","b")' and f.get("t") is None   # explícita, não compartilhada
        assert ed.style_of(r, 42) == ed.style_of(7, 42)
    assert ed.formula_at(5, 42) == '=IF(A5="x","a","b")'                 # filha do modelo: intacta
    assert report.orphan_formula_rows == 0 and report.overwritten_constants == 0


def test_shared_master_surrounded_by_constants_is_never_overwritten(tmp_path, editor):
    # AP3 e AP5 são constantes; o mestre compartilhado fica em AP4 e as filhas em AP6:AP7
    s, ed = editor("DESPESAS", desp_formula_rows=5, desp_constant_rows=(3, 5))
    master_xml = etree.tostring(ed._cell(4, 42, create=False))
    from papaki_dre.build import SheetReport, _extend_formulas
    report = SheetReport("DESPESAS", 3, 7, 5)
    for letter in ("AP", "AQ"):
        _extend_formulas(ed, "cp", "DESPESAS", letter, 3, 7, report)
    assert etree.tostring(ed._cell(4, 42, create=False)) == master_xml        # mestre byte a byte
    for r in range(3, 8):                                                      # constantes viram fórmula
        assert ed.formula_at(r, 42) == f'=IF(A{r}="x","a","b")'               # filhas 6 e 7 resolvem
        assert ed.formula_at(r, 43) == f'=LEN(AP{r})&"!"'
    assert report.overwritten_constants == 4                                   # AP e AQ nas linhas 3 e 5


def test_save_removes_calc_chain_and_forces_full_calc(tmp_path, editor):
    s, ed = editor()
    out = s.save(tmp_path / "o.xlsx")
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "xl/calcChain.xml" not in names
        assert b'fullCalcOnLoad="1"' in z.read("xl/workbook.xml")
        assert b"calcChain" not in z.read("[Content_Types].xml")
        assert b"calcChain" not in z.read("xl/_rels/workbook.xml.rels")


def test_all_other_parts_are_byte_identical(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx")
    with WorkbookSession(tpl) as s:
        s.editor("DESPESAS").write(3, 1, "x")
        out = s.save(tmp_path / "o.xlsx")
    diff = changed_parts(tpl, out)
    assert diff["changed"] == sorted(["xl/worksheets/sheet2.xml", "xl/workbook.xml",
                                      "[Content_Types].xml", "xl/_rels/workbook.xml.rels"])
    assert diff["removed"] == ["xl/calcChain.xml"]
    assert diff["added"] == []


def test_untouched_sheet_keeps_formula_and_cached_value(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx")
    with WorkbookSession(tpl) as s:
        s.editor("BASE").write(5, 1, "x")
        out = s.save(tmp_path / "o.xlsx")
    with _wb(out, data_only=True) as wb:
        assert wb["OUTRA"]["B1"].value == 10        # valor em cache da aba que não pode mudar


def test_dimension_grows_with_new_rows(tmp_path, editor):
    import re
    s, ed = editor("DESPESAS")
    ed.write(300, 1, "x")
    out = s.save(tmp_path / "o.xlsx")
    xml = _read_part(out, "xl/worksheets/sheet2.xml").decode()
    assert re.search(r'<dimension ref="A2:AQ300"', xml)


def _sha(path):
    import hashlib
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_failed_save_keeps_previous_output_and_leaves_no_temp_file(tmp_path):
    out = tmp_path / "o.xlsx"
    out.write_bytes(b"saida anterior")
    tpl = make_template(tmp_path / "t.xlsx", calc_pr=None)     # workbook.xml sem <calcPr>
    with WorkbookSession(tpl) as s:
        s.editor("BASE").write(5, 1, "x")
        with pytest.raises(ContractError, match="calcPr"):
            s.save(out)
    assert out.read_bytes() == b"saida anterior"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["o.xlsx", "t.xlsx"]


def test_failure_while_writing_the_zip_keeps_previous_output(tmp_path, monkeypatch):
    out = tmp_path / "o.xlsx"
    out.write_bytes(b"saida anterior")
    tpl = make_template(tmp_path / "t.xlsx")
    real = zipfile.ZipFile.writestr
    written = []

    def boom(self, zinfo_or_arcname, *a, **kw):
        name = getattr(zinfo_or_arcname, "filename", zinfo_or_arcname)
        if name == "xl/worksheets/sheet3.xml":            # a aba RECEITA, no meio do pacote
            raise OSError("disco cheio")
        written.append(name)
        return real(self, zinfo_or_arcname, *a, **kw)

    monkeypatch.setattr(zipfile.ZipFile, "writestr", boom)
    with WorkbookSession(tpl) as s:
        s.editor("BASE").write(5, 1, "x")
        with pytest.raises(OSError, match="disco cheio"):
            s.save(out)
    assert written                                         # falhou no meio da escrita, não antes
    assert out.read_bytes() == b"saida anterior"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["o.xlsx", "t.xlsx"]


@pytest.mark.parametrize("alias", [
    pytest.param(lambda d: d / "sub" / ".." / "t.xlsx", id="subpasta"),     # entra e volta
    pytest.param(lambda d: d / "sub" / ".." / "T.XLSX", id="subpasta-maiusculas",
                 marks=pytest.mark.skipif(sys.platform != "win32", reason="só no Windows")),
])
def test_saving_over_the_template_is_refused_and_template_is_intact(tmp_path, alias):
    tpl = make_template(tmp_path / "t.xlsx")
    (tmp_path / "sub").mkdir()
    other = alias(tmp_path)
    assert str(other) != str(tpl)
    before = _sha(tpl)
    with WorkbookSession(tpl) as s:
        s.editor("BASE").write(5, 1, "x")
        with pytest.raises(ContractError, match="modelo"):
            s.save(other)
    assert _sha(tpl) == before
    with WorkbookSession(tpl) as s:                 # continua um .xlsx válido e intocado
        assert s.editor("BASE").text_at(5, 1) is None


@pytest.mark.parametrize("calc_pr, expected", [
    ('<calcPr calcId="191029" fullCalcOnLoad="0"/>', '<calcPr calcId="191029" fullCalcOnLoad="1"/>'),
    ("<calcPr fullCalcOnLoad='0' calcId=\"191029\"/>", '<calcPr fullCalcOnLoad="1" calcId="191029"/>'),
    ('<calcPr fullCalcOnLoad="1" calcId="191029"/>', '<calcPr fullCalcOnLoad="1" calcId="191029"/>'),
    ('<calcPr calcId="191029"/>', '<calcPr fullCalcOnLoad="1" calcId="191029"/>'),
    ('<calcPr/>', '<calcPr fullCalcOnLoad="1"/>'),
    ('<calcPr calcId="1" fullCalcOnLoad = "false" iterate="1"></calcPr>',
     '<calcPr calcId="1" fullCalcOnLoad="1" iterate="1"></calcPr>'),
])
def test_full_calc_on_load_is_forced(tmp_path, calc_pr, expected):
    tpl = make_template(tmp_path / "t.xlsx", calc_pr=calc_pr)
    with WorkbookSession(tpl) as s:
        out = s.save(tmp_path / "o.xlsx")
    xml = _read_part(out, "xl/workbook.xml").decode()
    assert expected in xml and xml.count("fullCalcOnLoad") == 1


def test_full_calc_on_load_with_namespace_prefix():
    from papaki_dre.xlsx_inject import _force_full_calc
    assert (_force_full_calc(b'<x:workbook><x:calcPr calcId="1" fullCalcOnLoad="0"/></x:workbook>')
            == b'<x:workbook><x:calcPr calcId="1" fullCalcOnLoad="1"/></x:workbook>')
    assert (_force_full_calc(b'<x:workbook><x:calcPr calcId="1"/></x:workbook>')
            == b'<x:workbook><x:calcPr fullCalcOnLoad="1" calcId="1"/></x:workbook>')


from xlsx_factory import replace_part, sheet_xml


def test_missing_or_invalid_template_is_a_contract_error(tmp_path):
    with pytest.raises(ContractError, match="abrir o modelo"):
        WorkbookSession(tmp_path / "nao_existe.xlsx")
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"isto nao e um zip")
    with pytest.raises(ContractError, match="abrir o modelo"):
        WorkbookSession(bad)


def test_package_without_workbook_xml_is_a_contract_error_and_releases_the_file(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx")
    broken = replace_part(tpl, tmp_path / "b.xlsx", "xl/workbook.xml", b"", drop=True)
    with pytest.raises(ContractError, match="workbook.xml"):
        WorkbookSession(broken)
    broken.unlink()                    # no Windows só apaga se o construtor fechou o zip


def test_sheet_relationship_without_target_is_a_contract_error(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx")
    with zipfile.ZipFile(tpl) as z:
        rels = z.read("xl/_rels/workbook.xml.rels").replace(b'Id="rId3"', b'Id="rId99"')
    broken = replace_part(tpl, tmp_path / "b.xlsx", "xl/_rels/workbook.xml.rels", rels)
    with pytest.raises(ContractError, match="RECEITA"):
        WorkbookSession(broken)
    broken.unlink()


def _with_base(tmp_path, rows_xml):
    tpl = make_template(tmp_path / "t.xlsx")
    return replace_part(tpl, tmp_path / "b.xlsx", "xl/worksheets/sheet1.xml", sheet_xml(rows_xml))


def test_row_without_r_is_a_contract_error(tmp_path):
    with WorkbookSession(_with_base(tmp_path, '<row><c r="A1"><v>1</v></c></row>')) as s:
        with pytest.raises(ContractError, match="linha sem o atributo r"):
            s.editor("BASE")


def test_cell_without_r_is_a_contract_error(tmp_path):
    with WorkbookSession(_with_base(tmp_path, '<row r="1"><c><v>1</v></c></row>')) as s:
        with pytest.raises(ContractError, match="linha 1"):
            s.editor("BASE")


def test_cell_from_another_row_is_a_contract_error(tmp_path):
    with WorkbookSession(_with_base(tmp_path, '<row r="1"><c r="A2"><v>1</v></c></row>')) as s:
        with pytest.raises(ContractError, match="A2"):
            s.editor("BASE")


def test_shared_formula_child_without_master_is_a_contract_error(tmp_path):
    rows = '<row r="1"><c r="A1"><f t="shared" si="7"/><v>1</v></c></row>'
    with WorkbookSession(_with_base(tmp_path, rows)) as s:
        ed = s.editor("BASE")
        with pytest.raises(ContractError, match="A1"):
            ed.formula_at(1, 1)


def test_untranslatable_shared_formula_is_a_contract_error(tmp_path):
    # mestre em A5 aponta para B1; a filha em A3 cairia na linha -1
    rows = ('<row r="3"><c r="A3"><f t="shared" si="0"/></c></row>'
            '<row r="5"><c r="A5"><f t="shared" ref="A3:A5" si="0">B1</f></c></row>')
    with WorkbookSession(_with_base(tmp_path, rows)) as s:
        with pytest.raises(ContractError, match="A3"):
            s.editor("BASE").formula_at(3, 1)


def test_surrogates_and_noncharacters_are_removed_from_text(tmp_path, editor):
    s, ed = editor()
    ed.write(5, 1, "a\ud800b\udfffc\ufffed\uffffe\x1ff")
    s.save(tmp_path / "o.xlsx")
    with _wb(tmp_path / "o.xlsx") as wb:
        assert wb["BASE"]["A5"].value == "abcdef"


def test_formula_with_illegal_characters_is_a_contract_error(tmp_path, editor):
    s, ed = editor()
    with pytest.raises(ContractError, match="A5"):
        ed.write(5, 1, Formula('="a\ud800"'))


def test_timezone_aware_datetime_is_a_contract_error(tmp_path, editor):
    with pytest.raises(ContractError, match="fuso"):
        to_serial(dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc))
    s, ed = editor()
    with pytest.raises(ContractError, match="B5"):
        ed.write(5, 2, dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc))


def test_integers_beyond_excel_precision_are_a_contract_error(tmp_path, editor):
    s, ed = editor()
    ed.write(5, 1, 10 ** 15)                                    # ainda exato no Excel
    ed.write(6, 1, -(10 ** 15))
    with pytest.raises(ContractError, match="A7"):
        ed.write(7, 1, 10 ** 15 + 1)
    with pytest.raises(ContractError, match="A8"):
        ed.write(8, 1, -(10 ** 16))


def test_unknown_sheet_is_a_contract_error(tmp_path):
    with WorkbookSession(make_template(tmp_path / "t.xlsx")) as s:
        with pytest.raises(ContractError, match="NAO_EXISTE"):
            s.editor("NAO_EXISTE")


def test_shared_formula_children_reuse_one_translator_per_si(tmp_path, editor):
    s, ed = editor("DESPESAS", desp_formula_rows=5)
    assert [ed.formula_at(r, 42) for r in (4, 5, 6)] == [f'=IF(A{r}="x","a","b")' for r in (4, 5, 6)]
    assert list(ed._shared) == ["0"]


def test_formula_shifter_matches_translate_formula():
    from papaki_dre.xlsx_inject import FormulaShifter, translate_formula
    sh = FormulaShifter('=SUM($A$1:A2)+B2', "C2")
    for dest in ("C3", "D10", "C2"):
        assert sh.to(dest) == translate_formula('=SUM($A$1:A2)+B2', "C2", dest)
    with pytest.raises(ContractError, match="C1"):
        FormulaShifter("=A1", "A2").to("C1")


def test_last_formula_row_ignores_formulas_of_another_family(tmp_path, editor):
    s, ed = editor("DESPESAS", desp_formula_rows=5, desp_footer_row=9)
    assert ed.last_formula_row(42, first=3) == 7       # COUNTA(AP3:AP8) na 9 não é da família
    assert ed.last_formula_row(43, first=3) == 7       # $A$1&"x" na 9 também não
    assert ed.last_formula_row(1, first=3) is None


# ---- entradas malformadas viram ContractError (re-revisão 3) --------------------------

def _styles(numfmts="", xfs='<xf numFmtId="0"/>'):
    return (f'<styleSheet xmlns="{NS}"><numFmts>{numfmts}</numFmts>'
            f'<cellXfs>{xfs}</cellXfs></styleSheet>').encode()


@pytest.mark.parametrize("numfmt", ['<numFmt formatCode="dd/mm/yyyy"/>',
                                    '<numFmt numFmtId="abc" formatCode="dd/mm/yyyy"/>',
                                    '<numFmt numFmtId="1.5" formatCode="dd/mm/yyyy"/>'])
def test_numfmt_without_valid_id_is_a_contract_error(numfmt):
    with pytest.raises(ContractError, match="styles.xml.*numFmtId"):
        DateStyles(_styles(numfmts=numfmt))


def test_xf_with_invalid_numfmtid_is_a_contract_error():
    with pytest.raises(ContractError, match="styles.xml.*numFmtId"):
        DateStyles(_styles(xfs='<xf numFmtId="x"/>'))


def _sheet_editor(rows_xml, dim="A1:A1"):
    from xlsx_factory import sheet_xml
    return SheetEditor(sheet_xml(rows_xml, dim).encode("utf-8"), "ABA")


def test_cell_column_beyond_xfd_is_a_contract_error():
    with pytest.raises(ContractError, match="ABA.*AAAA1"):
        _sheet_editor('<row r="1"><c r="AAAA1"><v>1</v></c></row>')


@pytest.mark.parametrize("r", ["²", "١"])
def test_row_number_with_non_ascii_digits_is_a_contract_error(r):
    with pytest.raises(ContractError, match="ABA.*linha"):
        _sheet_editor(f'<row r="{r}"><c r="A1"><v>1</v></c></row>')


def test_cell_reference_with_non_ascii_digits_is_a_contract_error():
    with pytest.raises(ContractError, match="ABA.*linha 1"):
        _sheet_editor('<row r="1"><c r="A١"><v>1</v></c></row>')


@pytest.mark.parametrize("dim", ["ZZZZ1:A2", "lixo", "A1:AAAA9"])
def test_malformed_dimension_is_a_contract_error(dim):
    ed = _sheet_editor('<row r="1"><c r="A1"><v>1</v></c></row>', dim=dim)
    with pytest.raises(ContractError, match="ABA.*dimension"):
        ed.to_bytes()


@pytest.mark.parametrize("cells", [
    '<c r="B3"><f>"abc</f></c>',                                    # string sem fim: ilegível
    '<c r="B3"><f t="shared" si="9"/></c>',                        # filha sem mestre
    '<c r="B3"><f>A1)</f></c></row><row r="4"><c r="B4"><f>A1+#REF</f></c>',
])
def test_column_with_only_unreadable_formulas_says_so(cells):
    ed = _sheet_editor(f'<row r="3">{cells}</row>')
    with pytest.raises(ContractError, match=r"ABA!B.*ilegível.*intraduzível.*B3"):
        ed.last_formula_row(2, first=3)


def test_unreadable_formula_next_to_data_formulas_is_still_another_family(tmp_path):
    ed = _sheet_editor('<row r="3"><c r="B3"><f>A3*2</f></c></row>'
                       '<row r="4"><c r="B4"><f>A4*2</f></c></row>'
                       '<row r="9"><c r="B9"><f>"abc</f></c></row>')
    assert ed.last_formula_row(2, first=3) == 4


def test_invalid_value_leaves_no_cell_style_or_date_warning(tmp_path, editor):
    s, ed = editor("RECEITA")
    bad = dt.datetime(2026, 1, 5, tzinfo=dt.timezone.utc)
    with pytest.raises(ContractError):
        ed.write(50, 2, bad, style_from=(2, 2), date_styles=s.date_styles)
    assert not ed.row_exists(50)
    assert ed.date_warnings == {} and ed.touched_rows == set()
    # linha existente, coluna nova: a célula também não pode sobrar
    with pytest.raises(ContractError):
        ed.write(2, 30, float("nan"), style_from=(2, 2))
    assert ed._cell(2, 30, create=False) is None


@pytest.mark.parametrize("formula", ["=A1)", "=(A1))"])
def test_unbalanced_parentheses_are_a_contract_error_not_a_raw_exception(formula):
    from papaki_dre.xlsx_inject import FormulaShifter
    with pytest.raises(ContractError, match="ABA!B3"):
        FormulaShifter(formula, "B4", where="ABA").to("B3")


# ---- revisão final (rev4) ---------------------------------------------------------------

def _extend(ed, letter, first, last):
    from papaki_dre.build import SheetReport, _extend_formulas
    report = SheetReport("ABA", first, last, last - first + 1)
    _extend_formulas(ed, "cp", "ABA", letter, first, last, report)
    return report


def test_data_formula_reaching_above_first_has_no_reference_row():
    # limite documentado em last_formula_row: =A{r-3} só existe a partir da linha 6; nenhuma fórmula se
    # traduz para first=3 (cairia em A0), então não há base de extensão (na carga: classificação ausente)
    ed = _sheet_editor("".join(f'<row r="{r}"><c r="B{r}"><f>A{r - 3}</f></c></row>' for r in (6, 7, 8)))
    assert ed.last_formula_row(2, first=3) is None


def test_foreign_shared_master_inside_the_loaded_rows_is_a_contract_error():
    # SUM(A1:A4) na 5 é um resumo (traduzida para a 3 cairia em A-1); a filha B9 depende do mestre
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row><row r="4"><c r="B4"><f>A4*2</f></c></row>'
            '<row r="5"><c r="B5"><f t="shared" ref="B5:B9" si="0">SUM(A1:A4)</f></c></row>'
            '<row r="9"><c r="B9"><f t="shared" si="0"/></c></row>')
    ed = _sheet_editor(rows)
    with pytest.raises(ContractError, match="ABA!B5.*mestre"):
        _extend(ed, "B", 3, 6)


def test_relative_variant_of_the_data_formula_inside_the_loaded_rows_is_kept():
    # como DESPESAS!AP881:AP1070 do modelo de setembro: outra versão da classificação (com referências
    # relativas, em fórmula compartilhada) no meio das linhas carregadas não é resumo e fica intacta
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row><row r="4"><c r="B4"><f>A4*2</f></c></row>'
            '<row r="5"><c r="B5"><f t="shared" ref="B5:B6" si="0">A5*3</f></c></row>'
            '<row r="6"><c r="B6"><f t="shared" si="0"/></c></row>'
            '<row r="7"><c r="B7"><f>A7*2</f></c></row>')
    ed = _sheet_editor(rows)
    before = [etree.tostring(ed._cell(r, 2, create=False)) for r in range(3, 8)]
    report = _extend(ed, "B", 3, 8)
    assert [etree.tostring(ed._cell(r, 2, create=False)) for r in range(3, 8)] == before
    assert ed.formula_at(6, 2) == "=A6*3" and ed.formula_at(8, 2) == "=A8*2"
    assert report.overwritten_constants == 0


@pytest.mark.parametrize("cell, expected", [
    ("<f/>", True),                                 # vazia
    ('<f>"abc</f>', True),                          # ilegível
    ('<f t="shared" si="9"/>', True),               # filha sem mestre
    ("<f>COUNTA(B1:B4)</f>", True),                 # resumo relativo: sai da planilha em first
    ('<f>$A$1&amp;"x"</f>', True),                  # resumo só com referências absolutas
    ("<f>A5*2</f>", False),                         # da família
    ("<f>A5*3+A4</f>", False),                      # variante relativa
])
def test_is_foreign_formula(cell, expected):
    ed = _sheet_editor('<row r="3"><c r="B3"><f>A3*2</f></c></row>'
                       f'<row r="5"><c r="B5">{cell}</c><c r="C5"><v>1</v></c></row>')
    sig = ed.formula_family(2, first=3)[1]
    assert ed.is_foreign_formula(5, 2, 3, sig) is expected
    assert ed.is_foreign_formula(5, 3, 3, sig) is False      # sem <f>: constante, não fórmula


def test_cell_inside_an_array_formula_is_never_overwritten():
    # B5:B6 é uma fórmula de matriz; a filha B6 só tem valor e pareceria uma constante digitada
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row><row r="4"><c r="B4"><f>A4*2</f></c></row>'
            '<row r="5"><c r="B5"><f t="array" ref="B5:B6">A5:A6*2</f><v>1</v></c></row>'
            '<row r="6"><c r="B6"><v>2</v></c></row>')
    ed = _sheet_editor(rows)
    with pytest.raises(ContractError, match="ABA!B6.*matriz B5:B6"):
        ed.write(6, 2, 5)
    with pytest.raises(ContractError, match="matriz B5:B6"):
        _extend(ed, "B", 3, 6)
    assert ed.formula_at(5, 2) == "=A5:A6*2" and ed._cell(6, 2, create=False).findtext("{%s}v" % NS) == "2"


def test_string_that_is_empty_after_cleaning_is_treated_as_none(tmp_path):
    with WorkbookSession(make_template(tmp_path / "t.xlsx")) as s:
        ed = s.editor("BASE")
        ed.write(500, 1, "\x00\x01")
        ed.write(500, 2, "")
        assert not ed.row_exists(500) and ed.touched_rows == set()


def test_filter_name_with_several_areas_is_left_alone():
    from papaki_dre.xlsx_inject import _sync_filter_names
    wb = (b'<workbook><definedNames><definedName name="_xlnm._FilterDatabase" localSheetId="1" hidden="1">'
          b"DESPESAS!$A$2:$AR$9,DESPESAS!$AT$2:$AT$9</definedName></definedNames></workbook>")
    assert _sync_filter_names(wb, {1: "A2:AR500"}) == wb


def test_filter_name_with_namespace_prefix_is_extended():
    from papaki_dre.xlsx_inject import _sync_filter_names
    wb = (b'<x:workbook><x:definedNames><x:definedName name="_xlnm._FilterDatabase" localSheetId="1" '
          b'hidden="1">DESPESAS!$A$2:$AR$9</x:definedName></x:definedNames></x:workbook>')
    assert _sync_filter_names(wb, {1: "A2:AR500"}) == wb.replace(b"$AR$9", b"$AR$500")


def test_illegal_pattern_spells_the_noncharacters_as_escapes():
    import inspect
    from papaki_dre import xlsx_inject
    line = next(l for l in inspect.getsource(xlsx_inject).splitlines() if l.startswith("_ILLEGAL"))
    assert "\ufffe" not in line and "\uffff" not in line and r"\ufffe\uffff" in line


# ---- revisão 6 ----------------------------------------------------------------------------

def test_many_one_cell_arrays_do_not_make_each_write_scan_them_all(monkeypatch):
    # 3000 matrizes de 1 célula (C1:C3000) e 3000 gravações em A: antes cada gravação percorria a
    # lista inteira (O(gravações x matrizes), ~8,5 s). Agora o índice é por linha e ignora matrizes
    # de 1 célula; conta-se quantas entradas do índice são examinadas (nenhuma), não segundos.
    n = 3000
    ed = _sheet_editor("".join(f'<row r="{r}"><c r="C{r}"><f t="array" ref="C{r}">A{r}*2</f><v>0</v></c></row>'
                               for r in range(1, n + 1)))
    examined = 0
    real = SheetEditor.array_formula_at

    def counting(self, r, col):
        nonlocal examined
        examined += len(self._array_index().get(r, ()))
        return real(self, r, col)

    monkeypatch.setattr(SheetEditor, "array_formula_at", counting)
    for r in range(1, n + 1):
        ed.write(r, 1, r)
    assert examined == 0
    assert ed.formula_at(n, 3) == f"=A{n}*2"


def test_multi_cell_array_index_only_holds_the_rows_it_covers():
    ed = _sheet_editor('<row r="5"><c r="B5"><f t="array" ref="B5:C6">A5:A6*2</f></c></row>'
                       '<row r="9"><c r="B9"><f t="array" ref="B9">A9*2</f></c></row>')
    assert sorted(ed._array_index()) == [5, 6]
    assert ed.array_formula_at(6, 3) == "B5:C6" and ed.array_formula_at(6, 4) is None
    assert ed.array_formula_at(9, 2) is None            # 1 célula: a checagem normal de fórmula basta
    with pytest.raises(ContractError, match="ABA!B9 contém fórmula do modelo"):
        ed.write(9, 2, 1)


def test_one_cell_foreign_array_inside_the_loaded_rows_is_replaced_and_counted():
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row><row r="4"><c r="B4"><f>A4*2</f></c></row>'
            '<row r="5"><c r="B5"><f t="array" ref="B5">SUM(A1:A4)</f><v>1</v></c></row>')
    ed = _sheet_editor(rows)
    report = _extend(ed, "B", 3, 6)
    assert ed.formula_at(5, 2) == "=A5*2" and ed._cell(5, 2, create=False).find("{%s}f" % NS).get("t") is None
    assert ed.formula_at(6, 2) == "=A6*2"
    assert report.overwritten_constants == 1


def test_one_cell_array_of_the_family_is_kept():
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row>'
            '<row r="4"><c r="B4"><f t="array" ref="B4">A4*2</f><v>1</v></c></row>')
    ed = _sheet_editor(rows)
    before = etree.tostring(ed._cell(4, 2, create=False))
    report = _extend(ed, "B", 3, 5)
    assert etree.tostring(ed._cell(4, 2, create=False)) == before and report.overwritten_constants == 0


def test_foreign_anchor_of_a_multi_cell_array_is_refused_before_any_change():
    # o resumo SUM(A1:A4) é a âncora (não a filha) de uma matriz B5:B6; B4 tem uma constante que
    # seria substituída antes: nada muda e nada é contado
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row><row r="4"><c r="B4"><v>7</v></c></row>'
            '<row r="5"><c r="B5"><f t="array" ref="B5:B6">SUM(A1:A4)</f><v>1</v></c></row>')
    ed = _sheet_editor(rows)
    from papaki_dre.build import SheetReport, _extend_formulas
    report = SheetReport("ABA", 3, 5, 3)
    with pytest.raises(ContractError, match="ABA!B5 está dentro da fórmula de matriz B5:B6"):
        _extend_formulas(ed, "cp", "ABA", "B", 3, 5, report)
    assert ed.formula_at(4, 2) is None and ed._cell(4, 2, create=False).findtext("{%s}v" % NS) == "7"
    assert report.overwritten_constants == 0 and ed.touched_rows == set()


def test_foreign_shared_master_error_leaves_earlier_rows_untouched_and_uncounted():
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row><row r="4"><c r="B4"><v>7</v></c></row>'
            '<row r="5"><c r="B5"><f t="shared" ref="B5:B9" si="0">SUM(A1:A4)</f></c></row>'
            '<row r="9"><c r="B9"><f t="shared" si="0"/></c></row>')
    ed = _sheet_editor(rows)
    from papaki_dre.build import SheetReport, _extend_formulas
    report = SheetReport("ABA", 3, 6, 4)
    with pytest.raises(ContractError, match="ABA!B5.*mestre"):
        _extend_formulas(ed, "cp", "ABA", "B", 3, 6, report)
    assert ed._cell(4, 2, create=False).findtext("{%s}v" % NS) == "7" and ed.formula_at(4, 2) is None
    assert ed.formula_at(6, 2) is None
    assert report.overwritten_constants == 0 and ed.touched_rows == set()


def test_check_write_validates_without_changing_anything():
    ed = _sheet_editor('<row r="3"><c r="B3"><f>A3*2</f></c></row>')
    ed.check_write(4, 2, 5)
    assert not ed.row_exists(4)
    with pytest.raises(ContractError, match="ABA!B3 contém fórmula do modelo"):
        ed.check_write(3, 2, 5)
    with pytest.raises(ContractError, match="ABA!A5"):
        ed.check_write(5, 1, object())
    assert not ed.row_exists(5)


def test_formula_and_type_errors_of_write_name_the_sheet(tmp_path, editor):
    _, ed = editor("DESPESAS")
    with pytest.raises(ContractError, match="DESPESAS!AP3 contém fórmula do modelo"):
        ed.write(3, 42, "x")
    with pytest.raises(ContractError, match="tipo de valor não suportado em DESPESAS!A900"):
        ed.write(900, 1, object())


def test_signatures_are_computed_once_per_row(monkeypatch):
    # formula_family (que relê as linhas de baixo para cima) e is_foreign_formula (que, nas variantes,
    # ainda confere se a fórmula muda de uma linha para outra) sobre as mesmas linhas: cada fórmula é
    # tokenizada uma vez só (12 células com fórmula)
    import papaki_dre.xlsx_inject as xi
    ed = _sheet_editor("".join(f'<row r="{r}"><c r="B{r}"><f>A{r}*2</f></c></row>' for r in range(3, 13))
                       + "".join(f'<row r="{r}"><c r="B{r}"><f>A{r}*3</f></c></row>' for r in (13, 14)))
    built = 0
    real_init = xi.FormulaShifter.__init__

    def counting(self, *a, **kw):
        nonlocal built
        built += 1
        real_init(self, *a, **kw)

    monkeypatch.setattr(xi.FormulaShifter, "__init__", counting)
    sig = ed.formula_family(2, first=3)[1]
    for r in range(3, 15):
        assert ed.is_foreign_formula(r, 2, 3, sig) is False
    assert built == 12


def test_signature_cache_follows_a_rewritten_cell():
    ed = _sheet_editor('<row r="3"><c r="B3"><f>A3*2</f></c></row>'
                       '<row r="5"><c r="B5"><f>COUNTA(B1:B4)</f></c></row>')
    sig = ed.formula_family(2, first=3)[1]
    assert ed.is_foreign_formula(5, 2, 3, sig) is True
    ed.write(5, 2, Formula("=A5*2"), replace_formula=True)
    assert ed.is_foreign_formula(5, 2, 3, sig) is False


def test_shared_children_reuse_the_master_translator_for_their_signatures(monkeypatch):
    # B3 é o mestre (si=0) de B3:B12; as filhas não são retokenizadas para a assinatura: o tradutor
    # do mestre (o mesmo que formula_at usa) dá o mesmo resultado
    import papaki_dre.xlsx_inject as xi
    ed = _sheet_editor('<row r="3"><c r="B3"><f t="shared" ref="B3:B12" si="0">A3*2</f></c></row>'
                       + "".join(f'<row r="{r}"><c r="B{r}"><f t="shared" si="0"/></c></row>'
                                 for r in range(4, 13)))
    built = 0
    real_init = xi.FormulaShifter.__init__

    def counting(self, *a, **kw):
        nonlocal built
        built += 1
        real_init(self, *a, **kw)

    monkeypatch.setattr(xi.FormulaShifter, "__init__", counting)
    last, sig = ed.formula_family(2, first=3)
    assert (last, sig) == (12, "=A3*2")
    assert [ed.is_foreign_formula(r, 2, 3, sig) for r in range(3, 13)] == [False] * 10
    assert built == 2                                     # o do mestre como célula e o das filhas (si=0)


# ---- fórmula completa em toda a coluna: troca de mestres compartilhados (13/10/2026) -------------

_GROUP = ('<row r="3"><c r="B3"><f t="shared" ref="B3:B5" si="0">A3*2</f><v>1</v></c></row>'
          '<row r="4"><c r="B4"><f t="shared" si="0"/><v>1</v></c><c r="C4"><v>9</v></c></row>'
          '<row r="5"><c r="B5"><f t="shared" si="0"/></c></row>'
          '<row r="6"><c r="B6"><f t="shared" ref="B6:B7" si="1">A6*5</f></c></row>'
          '<row r="7"><c r="B7"><f t="shared" si="1"/></c></row>')


def test_write_batch_replaces_a_shared_master_when_the_whole_ref_is_in_the_batch():
    from xlsx_factory import orphan_shared_children
    ed = _sheet_editor(_GROUP)
    ed.write_batch([(r, 2, Formula(f"=A{r}*3")) for r in (3, 4, 5)])
    assert [ed.formula_at(r, 2) for r in range(3, 8)] == ["=A3*3", "=A4*3", "=A5*3", "=A6*5", "=A7*5"]
    assert ed._cell(3, 2, create=False).find("{%s}f" % NS).get("t") is None
    assert orphan_shared_children(ed.to_bytes()) == []
    assert ed._cell(4, 3, create=False).findtext("{%s}v" % NS) == "9"


def test_write_batch_refuses_a_master_whose_ref_is_not_covered_and_changes_nothing():
    ed = _sheet_editor(_GROUP)
    before = ed.to_bytes()
    with pytest.raises(ContractError, match=r"ABA!B3 é o mestre da fórmula compartilhada si=0 \(ref B3:B5\)"
                                            r".*ficariam de fora: B5"):
        ed.write_batch([(r, 2, Formula(f"=A{r}*3")) for r in (3, 4)])
    assert ed.to_bytes() == before and ed.touched_rows == set()


def test_write_batch_refuses_children_of_the_si_outside_the_ref():
    # filha B9 do si=0 fora do ref declarado B3:B5 (arquivo malformado): também ficaria sem mestre
    ed = _sheet_editor(_GROUP + '<row r="9"><c r="B9"><f t="shared" si="0"/></c></row>')
    with pytest.raises(ContractError, match=r"ficariam de fora: B9"):
        ed.write_batch([(r, 2, Formula(f"=A{r}*3")) for r in (3, 4, 5)])


def test_single_write_still_never_replaces_a_shared_master():
    ed = _sheet_editor(_GROUP)
    with pytest.raises(ContractError, match="ABA!B3 é o mestre de uma fórmula compartilhada"):
        ed.write(3, 2, Formula("=A3*3"), replace_formula=True)


def _unify(ed, letter, first, last, master_row):
    from papaki_dre.build import SheetReport, _extend_formulas
    report = SheetReport("ABA", first, last, last - first + 1)
    _extend_formulas(ed, "cp", "ABA", letter, first, last, report, master_row=master_row)
    return report


def test_illegible_configured_formula_is_a_contract_error():
    ed = _sheet_editor('<row r="3"><c r="B3"><f>A3*2</f></c></row><row r="4"><c r="B4"><f>"abc</f></c></row>')
    with pytest.raises(ContractError, match=r"ABA!B4: a fórmula da linha configurada .*ilegível"):
        _unify(ed, "B", 3, 5, 4)
    assert ed.touched_rows == set()


def test_unification_keeps_identical_rows_and_counts_replaced_formulas_apart_from_constants():
    rows = ('<row r="3"><c r="B3"><f>A3*2</f></c></row>'
            '<row r="4"><c r="B4"><v>7</v></c></row>'
            '<row r="5"><c r="B5"><f>A5*2+A4</f></c></row>'
            '<row r="6"><c r="B6"><f>A6*2+A5</f></c></row>')
    ed = _sheet_editor(rows)
    kept = etree.tostring(ed._cell(6, 2, create=False))
    report = _unify(ed, "B", 3, 7, 5)
    assert [ed.formula_at(r, 2) for r in range(3, 8)] == [f"=A{r}*2+A{r - 1}" for r in range(3, 8)]
    assert etree.tostring(ed._cell(6, 2, create=False)) == kept
    assert report.unified_formula_rows == 1 and report.overwritten_constants == 1   # B3 | B4
