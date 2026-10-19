import datetime as dt

import openpyxl
import pytest

from papaki_dre.config import R_BLANK, R_FOOTER, R_H, R_MODELO, R_PESQ
from papaki_dre.headers import load_headers
from papaki_dre.sanitize.comissoes import sanitize_comissoes
from papaki_dre.types import ContractError, Formula
from xlsxcache import inject_cache

H = load_headers()["comissoes"]
MODELO = "Untitled - Publicidade Padrão"


def _build(tmp_path, data_rows, header=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "BASE DE COMISSÕES POR JOB — REAL x ORÇADO"
    for j, name in enumerate(header or H, start=1):
        ws.cell(3, j, name)
    for i, row in enumerate(data_rows, start=4):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    p = tmp_path / "com.xlsx"
    wb.save(p)
    return p


def _line(r, modelo=MODELO, titulo="Job X", comissao="Diretor", forn="ACME", orcado=100, real=0):
    return [modelo, "#75131", dt.datetime(2026, 9, 22), "Fulano", "Cliente", "Agência", titulo,
            comissao, forn, orcado, real, f"=IF(K{r}=0,J{r},0)", f"=J{r}-K{r}",
            f"=IF(J{r}=0,0,(J{r}-K{r})/J{r})"]


def test_keeps_matching_rows_and_preserves_types(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4)]))
    assert s.kept == 1 and s.balanced()
    assert s.rows[0][2] == dt.datetime(2026, 9, 22)
    assert s.rows[0][9] == 100


def test_formulas_are_readdressed_to_destination_rows(tmp_path):
    rows = [_line(4, modelo="Plena - Publicidade Padrão"), _line(5), _line(6)]
    s = sanitize_comissoes(_build(tmp_path, rows))
    assert s.kept == 2
    first, second = s.rows
    assert isinstance(first[11], Formula)
    assert first[11] == "=IF(K2=0,J2,0)"            # linha de destino 2
    assert second[11] == "=IF(K3=0,J3,0)"
    assert second[12] == "=J3-K3"
    assert s.source_rows == [5, 6]


def test_db_rows_have_values_not_formulas(tmp_path):
    p = _build(tmp_path, [_line(4, orcado=100, real=0)])
    inject_cache(p, {"L4": 100, "M4": 100, "N4": 1})      # o que o Excel teria calculado
    s = sanitize_comissoes(p)
    assert s.db_rows[0][11:14] == [100, 100, 1]
    assert all(isinstance(v, Formula) for v in s.rows[0][11:14])
    assert not s.warnings


def test_excludes_other_models(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4, modelo="Vision - Publicidade Padrão ")]))
    assert s.kept == 0 and [e.reason for e in s.exclusions] == [R_MODELO]


def test_model_comparison_ignores_surrounding_spaces(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4, modelo=MODELO + " ")]))
    assert s.kept == 1


def test_excludes_pesquisa_e_tratamento(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4, titulo="x | PESQUISA E TRATAMENTO | y")]))
    assert [e.reason for e in s.exclusions] == [R_PESQ]


@pytest.mark.parametrize("opcao", ["Seguro Produção", "Taxa da produtora", "Comissão total 1",
                                   "Comissão total 2", "Impostos", "Taxa ADM"])
def test_excludes_the_six_column_h_options(tmp_path, opcao):
    s = sanitize_comissoes(_build(tmp_path, [_line(4, comissao=opcao)]))
    assert [e.reason for e in s.exclusions] == [R_H]


def test_new_column_h_options_are_preserved(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4, comissao="Nany"), _line(5, comissao="Opção Nova")]))
    assert s.kept == 2


def test_reason_precedence_model_then_pesquisa_then_h(tmp_path):
    row = _line(4, modelo="Outro", titulo="PESQUISA E TRATAMENTO", comissao="Impostos")
    s = sanitize_comissoes(_build(tmp_path, [row]))
    assert [e.reason for e in s.exclusions] == [R_MODELO]


def test_reason_precedence_pesquisa_before_h(tmp_path):
    row = _line(4, titulo="x PESQUISA E TRATAMENTO", comissao="Impostos")
    s = sanitize_comissoes(_build(tmp_path, [row]))
    assert [e.reason for e in s.exclusions] == [R_PESQ]


def test_blank_rows_are_technical(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4), [None] * 14, _line(6)]))
    assert [t.reason for t in s.technical] == [R_BLANK]


def _footer(first, last):
    """Formato do bruto real: coluna A vazia, rótulo em H e somas em J:N."""
    row = [None] * 14
    row[7] = "TOTAL GERAL"
    for j, col in enumerate("JKLM", start=9):
        row[j] = f"=SUM({col}{first}:{col}{last})"
    row[13] = f"=IF(J{last + 2}=0,0,(J{last + 2}-K{last + 2})/J{last + 2})"
    return row


def test_total_geral_footer_is_technical_not_model_exclusion(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4), _line(5), [None] * 14, _footer(4, 5)]))
    assert s.kept == 2
    assert [t.reason for t in s.technical] == [R_BLANK, R_FOOTER]
    assert not s.exclusions
    assert s.balanced()


def test_total_geral_outside_column_h_is_not_footer(tmp_path):
    row = _footer(4, 4)
    row[7], row[6] = None, "TOTAL GERAL"
    s = sanitize_comissoes(_build(tmp_path, [_line(4), row]))
    assert R_FOOTER not in [t.reason for t in s.technical]


def test_unexpected_cached_error_in_kept_row_blocks(tmp_path):
    p = _build(tmp_path, [_line(4)])
    inject_cache(p, {"L4": "#REF!", "M4": 100, "N4": 1})
    with pytest.raises(ContractError, match="#REF!"):
        sanitize_comissoes(p)


def test_cached_error_in_excluded_row_does_not_block(tmp_path):
    p = _build(tmp_path, [_line(4, modelo="Outro")])
    inject_cache(p, {"L4": "#REF!"})
    assert sanitize_comissoes(p).kept == 0


def test_expected_div0_in_kept_row_becomes_null(tmp_path):
    p = _build(tmp_path, [_line(4)])
    inject_cache(p, {"L4": 100, "M4": 100, "N4": "#DIV/0!"})
    s = sanitize_comissoes(p)
    assert s.db_rows[0][13] is None
    assert s.rows[0][13] == "=IF(J2=0,0,(J2-K2)/J2)"


def test_formula_without_cache_warns_and_is_null(tmp_path):
    s = sanitize_comissoes(_build(tmp_path, [_line(4)]))      # openpyxl não grava cache
    assert s.db_rows[0][11:14] == [None, None, None]
    assert any("cache" in w for w in s.warnings)


def test_filled_header_beyond_ncols_stops(tmp_path):
    p = _build(tmp_path, [_line(4)])
    wb = openpyxl.load_workbook(p)
    wb.active.cell(3, 20, "Coluna extra")
    wb.save(p)
    with pytest.raises(ContractError, match="extra|além"):
        sanitize_comissoes(p)


def test_empty_cells_beyond_ncols_are_fine(tmp_path):
    p = _build(tmp_path, [_line(4)])
    wb = openpyxl.load_workbook(p)
    wb.active.cell(3, 20, "   ")
    wb.active.cell(4, 30, None)
    wb.save(p)
    assert sanitize_comissoes(p).kept == 1


def test_header_mismatch_stops(tmp_path):
    bad = H[:-1] + ["Var X"]
    with pytest.raises(ContractError):
        sanitize_comissoes(_build(tmp_path, [_line(4)], header=bad))


@pytest.mark.real
def test_real_september_file(comissoes_real):
    s = sanitize_comissoes(comissoes_real)
    assert s.balanced()
    assert s.raw_count == 1222
    assert s.kept == 649
    assert sorted(t.reason for t in s.technical) == [R_BLANK, R_FOOTER]
    assert len([e for e in s.exclusions if e.reason == R_MODELO]) == 81
    assert all(isinstance(r[11], Formula) for r in s.rows)
