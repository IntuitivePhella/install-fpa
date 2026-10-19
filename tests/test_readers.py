import io
import re
import zipfile

import openpyxl
import pytest

from papaki_dre.readers import (_materialize, convert_xls_to_xlsx, open_xlsx, read_delimited,
                                xlsx_cache_path)
from papaki_dre.types import ContractError


def test_read_delimited_utf8_with_bom(tmp_path):
    p = tmp_path / "a.csv"
    p.write_bytes("Cód.;Valor\nx-1;R$ 1,00\n".encode("utf-8-sig"))
    enc, rows = read_delimited(p)
    assert enc == "utf-8-sig"
    assert rows == [["Cód.", "Valor"], ["x-1", "R$ 1,00"]]


def test_read_delimited_falls_back_to_cp1252(tmp_path):
    p = tmp_path / "b.csv"
    p.write_bytes("Emissão;Descrição\n05/01/2026;ação\n".encode("cp1252"))
    enc, rows = read_delimited(p)
    assert enc == "cp1252"
    assert rows[1] == ["05/01/2026", "ação"]


def test_read_delimited_keeps_quoted_semicolons_and_newlines(tmp_path):
    p = tmp_path / "c.csv"
    p.write_bytes('a;b\n"x;y";"linha1\nlinha2"\n'.encode("utf-8"))
    _, rows = read_delimited(p)
    assert rows[1] == ["x;y", "linha1\nlinha2"]


def _xlsx_with_bad_autofilter(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BASE COMISSÕES"
    ws.append(["a", "b"])
    ws.append([1, 2])
    ws.auto_filter.ref = "A1:B2"
    wb.save(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb'<autoFilter ref="[^"]*"', b"<autoFilter ref=\"'X'!A3:N4\"", data)
            zout.writestr(item, data)
    path.write_bytes(buf.getvalue())


def test_open_xlsx_repairs_invalid_autofilter(tmp_path):
    p = tmp_path / "bad.xlsx"
    _xlsx_with_bad_autofilter(p)
    with pytest.raises(ValueError):
        openpyxl.load_workbook(p)          # confirma que o problema é real
    ws = open_xlsx(p).worksheets[0]
    assert [c.value for c in ws[2]] == [1, 2]


def test_open_xlsx_keeps_formulas_or_values_on_request(tmp_path):
    p = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 2
    wb.active["B1"] = "=A1*3"
    wb.save(p)
    assert open_xlsx(p).active["B1"].value == "=A1*3"


@pytest.mark.excel
@pytest.mark.real
def test_convert_xls_to_xlsx_keeps_formulas(rra_real, tmp_path):
    out = convert_xls_to_xlsx(rra_real, tmp_path / "rra.xlsx")
    ws = openpyxl.load_workbook(out)["Plan1"]
    assert ws.max_column >= 49
    formulas = sum(1 for r in ws.iter_rows(min_row=4) for c in r
                   if isinstance(c.value, str) and c.value.startswith("="))
    assert formulas > 1000


# --- cache da conversão XLS -> XLSX (sem Excel) ------------------------------------------

def test_cache_path_depends_on_content_not_on_name(tmp_path):
    a, b, c = tmp_path / "a", tmp_path / "b", tmp_path / "c"
    for d in (a, b, c):
        d.mkdir()
    (a / "rel.xls").write_bytes(b"conteudo 1")
    (b / "rel.xls").write_bytes(b"conteudo 2")
    (c / "rel.xls").write_bytes(b"conteudo 1")
    dst = tmp_path / "work" / "rel.xlsx"
    pa, pb, pc = (xlsx_cache_path(d / "rel.xls", dst) for d in (a, b, c))
    assert pa != pb                       # mesmo nome, conteúdos diferentes: destinos diferentes
    assert pa == pc                       # mesmo conteúdo: reaproveita
    assert pa.parent == dst.parent and pa.suffix == ".xlsx"
    assert pa.name.startswith("rel-")


def test_materialize_never_reuses_a_partial_file(tmp_path):
    target = tmp_path / "work" / "rel-abc.xlsx"

    def broken(tmp):
        tmp.write_bytes(b"parcial")
        raise RuntimeError("Excel caiu no meio do SaveAs")

    with pytest.raises(RuntimeError, match="no meio"):
        _materialize(target, broken)
    assert not target.exists()
    assert list(target.parent.iterdir()) == []          # temporário removido

    calls = []

    def good(tmp):
        calls.append(tmp)
        assert tmp.parent == target.parent and tmp != target
        tmp.write_bytes(b"completo")

    assert _materialize(target, good) == target
    assert target.read_bytes() == b"completo"
    assert _materialize(target, good) == target           # agora reaproveita
    assert len(calls) == 1


def test_materialize_requires_the_output(tmp_path):
    with pytest.raises(ContractError):
        _materialize(tmp_path / "x.xlsx", lambda tmp: None)


class _FakeWorkbook:
    def __init__(self, log):
        self.log = log

    def SaveAs(self, path, FileFormat):
        self.log.append(("SaveAs", FileFormat))
        raise RuntimeError("falha original no SaveAs")

    def Close(self, save):
        self.log.append("Close")
        raise RuntimeError("falha no Close")


class _FakeExcel:
    def __init__(self, log):
        self.log = log
        outer = self

        class _Books:
            def Open(self, path, **kw):
                outer.log.append(("Open", kw))
                return _FakeWorkbook(outer.log)
        self.Workbooks = _Books()

    def __setattr__(self, name, value):
        if name in ("Visible", "DisplayAlerts"):
            self.log.append((name, value))
        object.__setattr__(self, name, value)

    def Quit(self):
        self.log.append("Quit")
        raise RuntimeError("falha no Quit")


def test_excel_is_always_quit_without_masking_the_original_error(tmp_path, monkeypatch):
    win32 = pytest.importorskip("win32com.client")
    log = []
    monkeypatch.setattr(win32, "DispatchEx", lambda name: _FakeExcel(log))
    src = tmp_path / "rel.xls"
    src.write_bytes(b"x")
    with pytest.raises(RuntimeError, match="falha original no SaveAs"):
        convert_xls_to_xlsx(src, tmp_path / "work" / "rel.xlsx")
    opened = [e for e in log if isinstance(e, tuple) and e[0] == "Open"]
    assert opened and opened[0][1].get("UpdateLinks") == 0 and opened[0][1].get("ReadOnly")
    assert "Close" in log and log[-1] == "Quit"
    assert not any((tmp_path / "work").glob("*.xlsx"))


def test_excel_is_quit_even_if_setup_fails(tmp_path, monkeypatch):
    win32 = pytest.importorskip("win32com.client")
    log = []

    class _Bad(_FakeExcel):
        def __setattr__(self, name, value):
            if name == "Visible":
                raise RuntimeError("falha ao configurar")
            super().__setattr__(name, value)

    monkeypatch.setattr(win32, "DispatchEx", lambda name: _Bad(log))
    src = tmp_path / "rel.xls"
    src.write_bytes(b"y")
    with pytest.raises(RuntimeError, match="falha ao configurar"):
        convert_xls_to_xlsx(src, tmp_path / "work" / "rel.xlsx")
    assert log[-1] == "Quit"
