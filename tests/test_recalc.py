import hashlib
import tempfile

import pytest

from papaki_dre import recalc
from papaki_dre.recalc import XL_ERR_NA, _count_na, contar_nd
from papaki_dre.types import ContractError


def test_count_na_counts_only_the_na_code():
    values = ((XL_ERR_NA, 1.0), ("x", XL_ERR_NA), (-2146826281, None), (True, 0))
    assert _count_na(values) == 2


def test_count_na_handles_single_cell_and_empty():
    assert _count_na(XL_ERR_NA) == 1
    assert _count_na(None) == 0
    assert _count_na(()) == 0
    assert _count_na(5.0) == 0


def test_contar_nd_empty_range_returns_zero_without_excel(monkeypatch):
    def boom(*a, **k):
        raise AssertionError("não deveria abrir o Excel")
    monkeypatch.setattr(recalc.shutil, "copyfile", boom)
    assert contar_nd("nao_existe.xlsx", "DESPESAS", ("AP", "AQ"), 3, 2) == 0


def test_contar_nd_rejects_first_row_below_one():
    with pytest.raises(ValueError):
        contar_nd("nao_existe.xlsx", "DESPESAS", ("AP", "AQ"), 0, 5)


def _xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "DESPESAS"
    ws["AP3"], ws["AQ3"] = "=NA()", "=1+1"
    ws["AP4"], ws["AQ4"] = "=2+2", "=NA()"
    ws["AP5"], ws["AQ5"] = "=NA()", "=NA()"
    ws["AP9"] = "=NA()"                                  # fora da faixa consultada
    src = tmp_path / "gerado.xlsx"
    wb.save(src)
    return src


@pytest.mark.excel
def test_contar_nd_missing_sheet_raises_contract_error_and_cleans_up(tmp_path, monkeypatch):
    src = _xlsx(tmp_path)
    tmp_root = tmp_path / "tmproot"
    tmp_root.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_root))
    with pytest.raises(ContractError, match="ausente no arquivo"):
        contar_nd(src, "NAO_EXISTE", ("AP", "AQ"), 3, 5)
    assert list(tmp_root.iterdir()) == []


@pytest.mark.excel
def test_contar_nd_in_real_excel_leaves_original_untouched(tmp_path, monkeypatch):
    src = _xlsx(tmp_path)
    before = hashlib.sha256(src.read_bytes()).hexdigest()
    tmp_root = tmp_path / "tmproot"
    tmp_root.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_root))

    assert contar_nd(src, "DESPESAS", ("AP", "AQ"), 3, 5) == 4

    assert hashlib.sha256(src.read_bytes()).hexdigest() == before
    assert list(tmp_root.iterdir()) == []                # cópia temporária apagada
