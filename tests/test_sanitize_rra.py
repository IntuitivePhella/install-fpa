import datetime as dt

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from papaki_dre.config import R_BLANK, R_FOOTER, R_NUCLEO, R_PESQ
from papaki_dre.headers import load_headers
from papaki_dre.produtoras import regras_padrao
from papaki_dre.sanitize.rra import sanitize_rra, sanitize_rra_xlsx
from papaki_dre.types import ContractError, Formula
from xlsxcache import inject_cache

NUCLEO_RRA = regras_padrao().nucleo_rra
H = load_headers()["rra"]
I_JOB, I_NUCLEO, I_TITULO = 0, H.index("Núcleo"), H.index("Título do job")
I_FAT = H.index("VALOR FATURADO")


def _job(r, job="75131", nucleo=NUCLEO_RRA, titulo="Job X", fat=100.0, formula=True):
    row = [None] * len(H)
    row[I_JOB], row[1] = job, dt.datetime(2026, 3, 17)
    row[I_NUCLEO], row[I_TITULO], row[I_FAT] = nucleo, titulo, fat
    if formula:
        row[I_FAT + 1] = f"={get_column_letter(I_FAT + 1)}{r}*2"      # coluna ao lado do VALOR FATURADO
    return row


def _build(tmp_path, jobs, total=True, header=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan1"
    ws["A1"] = "RENDIMENTOS DOS JOBS"
    for j, name in enumerate(header or H, start=1):
        ws.cell(3, j, name)
    r = 5
    for row in jobs:
        for j, v in enumerate(row, start=1):
            ws.cell(r, j, v)
        r += 1
    if total:
        ws.cell(r, 1, "TOTAL")
        col = get_column_letter(I_FAT + 1)
        ws[f"{col}{r}"] = f"=SUM({col}5:{col}{r - 1})"
    p = tmp_path / "rra.xlsx"
    wb.save(p)
    return p


def test_keeps_only_untitled_nucleo(tmp_path):
    jobs = [_job(5), _job(6, job="1", nucleo="Projetos - Kaus"), _job(7, job="2")]
    s = sanitize_rra_xlsx(_build(tmp_path, jobs))
    assert s.kept == 2
    assert [e.reason for e in s.exclusions] == [R_NUCLEO]
    assert s.balanced()


def test_excludes_pesquisa_e_tratamento(tmp_path):
    s = sanitize_rra_xlsx(_build(tmp_path, [_job(5, titulo="x PESQUISA E TRATAMENTO y")]))
    assert s.kept == 0 and [e.reason for e in s.exclusions] == [R_PESQ]


def test_total_row_is_technical(tmp_path):
    s = sanitize_rra_xlsx(_build(tmp_path, [_job(5, fat=100.0), _job(6, fat=50.0)]))
    assert [t.reason for t in s.technical] == [R_BLANK, R_FOOTER]
    assert s.kept == 2 and s.balanced()
    assert not hasattr(s, "controls")


def test_formulas_are_readdressed_after_exclusions(tmp_path):
    jobs = [_job(5, job="1", nucleo="Projetos - Kaus"), _job(6, job="2"), _job(7, job="3")]
    s = sanitize_rra_xlsx(_build(tmp_path, jobs))
    col = get_column_letter(I_FAT + 1)
    f = s.rows[0][I_FAT + 1]
    assert isinstance(f, Formula)
    assert f == f"={col}5*2"                 # primeira linha mantida vai para a linha 5
    assert s.rows[1][I_FAT + 1] == f"={col}6*2"


def test_expected_div0_is_preserved_in_db_as_null(tmp_path):
    row = _job(5, formula=False)
    row[I_FAT + 1] = "=1/0"
    p = _build(tmp_path, [row])
    inject_cache(p, {f"{get_column_letter(I_FAT + 2)}5": "#DIV/0!"})   # cache t="e", como o Excel
    s = sanitize_rra_xlsx(p)                                          # não bloqueia
    assert s.kept == 1
    assert s.rows[0][I_FAT + 1] == "=1/0"
    assert s.db_rows[0][I_FAT + 1] is None
    assert not s.warnings


def test_unexpected_error_in_kept_row_blocks(tmp_path):
    row = _job(5, formula=False)
    ws_path = _build(tmp_path, [row])
    wb = openpyxl.load_workbook(ws_path)
    cell = wb["Plan1"].cell(5, I_FAT + 3)
    cell.value = "#REF!"
    cell.data_type = "e"
    wb.save(ws_path)
    with pytest.raises(ContractError, match="#REF!"):
        sanitize_rra_xlsx(ws_path)


def test_error_in_excluded_row_does_not_block(tmp_path):
    row = _job(5, nucleo="Projetos - Kaus", formula=False)
    p = _build(tmp_path, [row])
    wb = openpyxl.load_workbook(p)
    cell = wb["Plan1"].cell(5, I_FAT + 3)
    cell.value = "#REF!"
    cell.data_type = "e"
    wb.save(p)
    assert sanitize_rra_xlsx(p).kept == 0


def test_blank_row_is_technical(tmp_path):
    s = sanitize_rra_xlsx(_build(tmp_path, [_job(5)]))
    assert R_BLANK in [t.reason for t in s.technical]      # a linha 4 é vazia


def test_title_is_kept(tmp_path):
    assert sanitize_rra_xlsx(_build(tmp_path, [_job(5)])).title == "RENDIMENTOS DOS JOBS"


def test_unexpected_cached_error_in_formula_blocks(tmp_path):
    p = _build(tmp_path, [_job(5)])
    inject_cache(p, {f"{get_column_letter(I_FAT + 2)}5": "#N/A"})
    with pytest.raises(ContractError, match="#N/A"):
        sanitize_rra_xlsx(p)


def test_formula_without_cache_warns_and_is_null(tmp_path):
    s = sanitize_rra_xlsx(_build(tmp_path, [_job(5)]))
    assert s.db_rows[0][I_FAT + 1] is None
    assert any("cache" in w for w in s.warnings)


def test_filled_header_beyond_ncols_stops(tmp_path):
    p = _build(tmp_path, [_job(5)])
    wb = openpyxl.load_workbook(p)
    wb["Plan1"].cell(3, len(H) + 5, "Coluna extra")
    wb.save(p)
    with pytest.raises(ContractError, match="extra|além"):
        sanitize_rra_xlsx(p)


def test_empty_cells_beyond_ncols_are_fine(tmp_path):
    p = _build(tmp_path, [_job(5)])
    wb = openpyxl.load_workbook(p)
    wb["Plan1"].cell(3, 72, "")
    wb["Plan1"].cell(6, 72, None)
    wb.save(p)
    assert sanitize_rra_xlsx(p).kept == 1


def test_header_mismatch_stops(tmp_path):
    bad = H[:-1] + ["Coluna Nova"]
    with pytest.raises(ContractError):
        sanitize_rra_xlsx(_build(tmp_path, [_job(5)], header=bad))


@pytest.mark.excel
@pytest.mark.real
def test_real_september_file(rra_real, tmp_path):
    s = sanitize_rra(rra_real, work_dir=tmp_path)
    assert s.balanced()
    assert s.kept == 72
    assert len([e for e in s.exclusions if e.reason == R_PESQ]) == 9
    assert len(s.exclusions) == (1410 - 81) + 9
    assert s.raw_count == 2408
    assert len(s.technical) == 998
    assert [t.reason for t in s.technical].count(R_FOOTER) == 1
    assert not s.warnings
    assert any(isinstance(v, Formula) for v in s.rows[0])
