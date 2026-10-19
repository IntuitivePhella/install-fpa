"""Ponta a ponta do comando `importar` com os brutos REAIS de amostras/2026-09 e o banco falso em memória
(nada vai ao Supabase). Usa o Excel real (COM, instância própria) para converter o XLS do RRA e contar
#N/D. Não imprime nem grava linhas de dados reais."""
import copy
import hashlib
import json
import re
import zipfile

import openpyxl
import pytest

from papaki_dre import cli
from papaki_dre.config import R_BLANK, R_FOOTER, R_H, R_JOB, R_LOC, R_MODELO, R_NUCLEO, R_PESQ
from papaki_dre.recalc import contar_nd
from papaki_dre.xlsx_inject import WorkbookSession, changed_parts
from conftest import BRUTO, REAL_FILES, TEMPLATE
from fake_supabase import FakeClient

pytestmark = [pytest.mark.real, pytest.mark.excel]

# Contagens finais do Plano 1 (extração de 08/10/2026): bruto = mantidas + técnicas + exclusões.
# Outra extração muda os números: não ajuste para passar sem entender a diferença.
EXPECTED = {
    "rra": {"bruto": 2408, "tecnicas": 998, "mantidas": 72, "excluidas": {R_NUCLEO: 1329, R_PESQ: 9}},
    "cp": {"bruto": 36999, "tecnicas": 2, "mantidas": 1110},          # exclusões: 21.505 + 14.382 (R_JOB/R_LOC)
    "cr": {"bruto": 486, "tecnicas": 2, "mantidas": 484, "excluidas": {}},
    "comissoes": {"bruto": 1222, "tecnicas": 2, "mantidas": 649,
                  "excluidas": {R_MODELO: 81, R_PESQ: 126, R_H: 364}},
}
BASE_SHEETS = ("BASE", "DESPESAS", "RECEITA", "REAL X ORÇADO COMISSÕES")
FINAL_NAME = "Modelo DRE - Untitled - 09_2026.xlsx"
ND_ESPERADO = 231        # decisão 11: #N/D em DESPESAS!AP:AQ, linhas 3-1112

PAPAKI = BRUTO.parent / "bruto_papaki"
FILES_2 = {
    "cp": PAPAKI / "Contas a Pagar - Consolidado - 09_2026.csv",
    "cr": PAPAKI / "Contas a Receber - Untitled - 09_2026.csv",
    "rra": PAPAKI / "RRA - Consolidado - 09_2026.xls",
    "comissoes": PAPAKI / "Real x Orcado Comissoes - Untitled - 09_2026.xlsx",
}


def _argv(out, files=REAL_FILES, *extra):
    return ["importar", "--produtora", "untitled", "--competencia", "2026-09",
            "--cp", str(files["cp"]), "--cr", str(files["cr"]),
            "--rra", str(files["rra"]), "--comissoes", str(files["comissoes"]),
            "--modelo", str(TEMPLATE), "--saida", str(out),
            "--tela", "cp=37030,cr=483,rra=1410", *extra]


def _need(files):
    if not all(p.exists() for p in [TEMPLATE, *files.values()]):
        pytest.skip("modelo ou brutos de referência ausentes")


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    _need(REAL_FILES)
    out = tmp_path_factory.mktemp("e2e") / "saida"
    client = FakeClient()
    code = cli.main(_argv(out), client=client)
    return code, out, client


@pytest.fixture(scope="module")
def run2(run, tmp_path_factory):
    """2ª extração (bruto_papaki), mesma competência, no MESMO banco falso, em outra pasta."""
    _need(FILES_2)
    _, _, client0 = run
    client = copy.deepcopy(client0)      # não altera o estado compartilhado da 1ª rodada
    out = tmp_path_factory.mktemp("e2e2") / "saida"
    code = cli.main(_argv(out, FILES_2), client=client)
    return code, out, client


def _conf(out):
    return json.loads((out / "conferencia.json").read_text(encoding="utf-8"))


def _nd_do_alerta(conf):
    achados = [int(m.group(1)) for a in conf["alertas"] if (m := re.match(r"(\d+) células com #N/D", a))]
    assert len(achados) == 1, "esperava exatamente um alerta de #N/D"
    return achados[0]


def test_run_is_published_with_alerts_only(run):
    code, out, _ = run
    conf = _conf(out)
    assert code == 0, (code, conf["bloqueios"])
    assert conf["status"] == "com_alertas" and conf["bloqueios"] == []


def test_september_count_flow_closes(run):
    _, out, _ = run
    rel = _conf(out)["relatorios"]
    for key, exp in EXPECTED.items():
        r = rel[key]
        assert r["registros_bruto"] == exp["bruto"], key
        assert r["removidas_tecnicas"] == exp["tecnicas"], key
        assert r["mantidas"] == r["carregadas"] == exp["mantidas"], key
        assert r["registros_bruto"] == r["mantidas"] + r["removidas_tecnicas"] + sum(r["excluidas"].values()), key
        if "excluidas" in exp:
            assert r["excluidas"] == exp["excluidas"], key
    assert set(rel["cp"]["excluidas"]) == {R_JOB, R_LOC}
    assert sorted(rel["cp"]["excluidas"].values()) == [14382, 21505]
    assert sorted(rel["cp"]["tecnicas_por_motivo"]) == sorted([R_BLANK, R_FOOTER])
    assert {k: v["carregadas"] for k, v in rel.items()} == {
        "rra": 72, "cp": 1110, "cr": 484, "comissoes": 649}


def test_screen_count_differences_are_alerts_only(run):
    _, out, _ = run
    conf = _conf(out)
    alertas = "\n".join(conf["alertas"])
    assert "tela mostrava 37030" in alertas and "tem 36997" in alertas      # CP: risco 3 da spec
    assert "tela mostrava 483" in alertas and "tem 484" in alertas          # CR
    assert not any("RRA" in a and "tela" in a for a in conf["alertas"])     # RRA: 2408 - 998 = 1410
    assert not any("tela" in b for b in conf["bloqueios"])
    # órfãs de DESPESAS/RECEITA (Plano 2): só alerta
    assert any("DESPESAS" in a and "47 linhas com fórmulas" in a for a in conf["alertas"])
    assert any("RECEITA" in a and "179 linhas com fórmulas" in a for a in conf["alertas"])


def test_nd_alert_counts_classification_without_match(run):
    code, out, _ = run
    conf = _conf(out)
    n = _nd_do_alerta(conf)
    # estável entre duas leituras independentes do arquivo entregue (cada uma em instância própria do Excel)
    final = out / FINAL_NAME
    lido1 = contar_nd(final, "DESPESAS", ("AP", "AQ"), 3, 1112)
    lido2 = contar_nd(final, "DESPESAS", ("AP", "AQ"), 3, 1112)
    assert lido1 == lido2 == n
    assert n == ND_ESPERADO
    assert not any("recalcular" in a for a in conf["alertas"])
    assert "contagem de #N/D não executada" not in conf["informacoes"]


def test_informative_fields_in_report(run):
    _, out, _ = run
    conf = _conf(out)
    d = conf["abas"]["DESPESAS"]
    assert d["constantes_substituidas"] == 439 and d["formulas_unificadas"] == 846
    assert isinstance(d["autofiltro_estendido"], bool)
    for name, a in conf["abas"].items():
        assert isinstance(a["autofiltro_estendido"], bool), name
        assert a["sem_formula"] == 0, name
    # só informação: nada disso vira alerta nem bloqueio
    for texto in conf["alertas"] + conf["bloqueios"]:
        assert "substituíd" not in texto and "unificada" not in texto and "autofiltro" not in texto
    assert any("439" in i for i in conf["informacoes"]) and any("846" in i for i in conf["informacoes"])
    md = (out / "conferencia.md").read_text(encoding="utf-8")
    assert "## Informações" in md and "439" in md and "846" in md


def test_database_received_one_current_snapshot(run):
    _, out, client = run
    imps = client.tables["importacoes"]
    assert len(imps) == 1
    assert imps[0]["status"] == "concluida" and imps[0]["vigente"] is True
    assert len(client.tables["rra_base"]) == 72
    assert len(client.tables["cp_despesas"]) == 1110
    assert len(client.tables["cr_receita"]) == 484
    assert len(client.tables["comissoes"]) == 649
    by_type = {}
    for e in client.tables["exclusoes"]:
        by_type[e["tipo"]] = by_type.get(e["tipo"], 0) + 1
    tecnicas = sum(EXPECTED[k]["tecnicas"] for k in EXPECTED)                 # 998 + 2 + 2 + 2
    negocio = sum(EXPECTED[k]["bruto"] - EXPECTED[k]["tecnicas"] - EXPECTED[k]["mantidas"] for k in EXPECTED)
    assert by_type == {"tecnica": 998 + 2 + 2 + 2, "negocio": 1338 + 35887 + 0 + 571}
    assert by_type == {"tecnica": tecnicas, "negocio": negocio}
    # cargas: 4 abas, todas com o hash do arquivo entregue
    sha = hashlib.sha256((out / FINAL_NAME).read_bytes()).hexdigest()
    cargas = client.tables["cargas_excel"]
    assert len(cargas) == 4
    assert {c["aba"] for c in cargas} == set(BASE_SHEETS)
    assert {c["sha256_saida"] for c in cargas} == {sha}
    assert {c["arquivo_saida"] for c in cargas} == {FINAL_NAME}
    assert {c["aba"]: c["linhas"] for c in cargas} == {
        "BASE": 72, "DESPESAS": 1110, "RECEITA": 484, "REAL X ORÇADO COMISSÕES": 649}


def test_output_workbook_has_the_data_and_no_leftovers(run):
    _, out, _ = run
    assert (out / FINAL_NAME).is_file()
    assert [p.name for p in out.glob("Modelo DRE*.xlsx")] == [FINAL_NAME]
    assert list(out.glob("*.novo.xlsx")) == [] and list(out.glob(".*.novo.xlsx")) == []
    ws = openpyxl.load_workbook(out / FINAL_NAME)["DESPESAS"]
    assert ws["A3"].value is not None and ws["A1112"].value is not None
    assert ws["A1113"].value is None, "A1113 deveria estar vazia"


def test_other_sheets_are_untouched(run):
    _, out, _ = run
    conf = _conf(out)
    with WorkbookSession(out / FINAL_NAME) as sess:
        base_parts = {sess.parts[n] for n in BASE_SHEETS}
    changed = set(conf["partes_alteradas"]["changed"])
    assert changed == base_parts | {"xl/workbook.xml", "[Content_Types].xml", "xl/_rels/workbook.xml.rels"}
    assert conf["partes_alteradas"]["removed"] == ["xl/calcChain.xml"]
    assert conf["partes_alteradas"]["added"] == []
    # conferência independente: todas as demais worksheets idênticas byte a byte ao modelo
    assert changed_parts(TEMPLATE, out / FINAL_NAME) == conf["partes_alteradas"]
    with zipfile.ZipFile(TEMPLATE) as a, zipfile.ZipFile(out / FINAL_NAME) as b:
        sheets = [n for n in a.namelist() if n.startswith("xl/worksheets/") and n.endswith(".xml")]
        assert len(sheets) == 17
        for n in sheets:
            if n not in base_parts:
                assert hashlib.sha256(a.read(n)).hexdigest() == hashlib.sha256(b.read(n)).hexdigest(), n


def test_second_run_with_same_files_is_ignored(run, tmp_path):
    _, _, client = run
    assert client.tables["importacoes"][0]["status"] == "concluida"
    n_linhas = len(client.tables["cp_despesas"])
    assert cli.main(_argv(tmp_path / "again"), client=client) == 0
    assert len(client.tables["importacoes"]) == 1
    assert len(client.tables["cp_despesas"]) == n_linhas
    assert not (tmp_path / "again").exists()           # idêntico: nenhum arquivo gerado


def test_second_extraction_is_not_ignored(run2, capsys):
    code, out, client = run2
    conf = _conf(out)
    imps = client.tables["importacoes"]
    # contagens observadas por relatório (só números), úteis para o Plano 3 / README
    rel = {k: {c: r[c] for c in ("registros_bruto", "removidas_tecnicas", "mantidas", "carregadas")}
           | {"excluidas": r["excluidas"]} for k, r in conf["relatorios"].items()}
    with capsys.disabled():
        print("\n2a extracao codigo:", code, "status:", conf["status"], "contagens:", rel,
              "bloqueios:", conf["bloqueios"], "alertas:", len(conf["alertas"]))
    assert code == 0, conf["bloqueios"]
    assert conf["status"] in ("ok", "com_alertas") and conf["bloqueios"] == []
    assert len(imps) == 2                               # hashes diferentes: não foi ignorada
    assert [(i["status"], i["vigente"]) for i in imps] == [("concluida", False), ("concluida", True)]
    id1, id2 = imps[0]["id"], imps[1]["id"]

    def n(table, imp_id, **filt):
        return sum(1 for r in client.tables[table]
                   if r["importacao_id"] == imp_id and all(r.get(k) == v for k, v in filt.items()))

    # a 1ª foto continua intacta; a nova tem as suas próprias linhas
    for table, c1, c2 in (("rra_base", 72, 72), ("cp_despesas", 1110, 1101),
                          ("cr_receita", 484, 481), ("comissoes", 649, 649)):
        assert (n(table, id1), n(table, id2)) == (c1, c2), table
    assert n("exclusoes", id1) == 1004 + 37796
    assert n("exclusoes", id2, tipo="tecnica") == 998 + 2 + 2 + 2
    assert n("exclusoes", id2, tipo="negocio") == (1325 + 9) + (21225 + 14318) + 0 + (364 + 81 + 126)
