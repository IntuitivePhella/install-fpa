"""Edição cirúrgica de um .xlsx: só o XML das abas pedidas muda; o resto do pacote
(tabelas dinâmicas, gráficos, estilos, strings compartilhadas) é copiado byte a byte."""
from __future__ import annotations

import bisect
import datetime as dt
import math
import os
import re
import tempfile
import unicodedata
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

from .types import ContractError, Formula, XlError

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
_EPOCH = dt.datetime(1899, 12, 30)
# Caracteres que o XML 1.0 não aceita: controles, surrogates soltos e as não-caracteres U+FFFE/U+FFFF.
_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_REF = re.compile(r"^([A-Z]+)([0-9]+)$")          # só dígitos ASCII (\d aceitaria "١")
_ROWNUM = re.compile(r"^[0-9]+$")
# Maior inteiro gravado como número: o Excel guarda 15 dígitos significativos (10**15 ainda é exato).
MAX_EXACT_INT = 10 ** 15


def _q(tag: str) -> str:
    return f"{{{NS}}}{tag}"


def to_serial(value) -> float | int:
    """datetime/date -> número de série do Excel (sistema 1900).

    Datas com fuso (tzinfo) são recusadas com ContractError: o Excel não guarda fuso, e converter
    para "o horário local" seria ambíguo (o da máquina? o de quem gerou o bruto?)."""
    if isinstance(value, dt.datetime) and value.tzinfo is not None:
        raise ContractError(f"data com fuso horário não é suportada pelo Excel: {value.isoformat()}")
    if not isinstance(value, dt.datetime):
        value = dt.datetime(value.year, value.month, value.day)
    delta = value - _EPOCH
    serial = delta.days + delta.seconds / 86400 + delta.microseconds / 86400e6
    return int(serial) if serial == int(serial) else serial


def _num(x, ref: str = "?") -> str:
    if isinstance(x, int):
        if abs(x) > MAX_EXACT_INT:
            raise ContractError(f"{ref}: inteiro {x} excede a precisão do Excel (15 dígitos)")
        return str(x)
    if not math.isfinite(x):
        raise ContractError(f"{ref}: número inválido para o Excel: {x!r}")
    if x == int(x) and abs(x) < 1e15:
        return str(int(x))
    return repr(x)


# O tokenizer do openpyxl estoura IndexError em parênteses desbalanceados (ex.: "=A1)").
_TRANSLATE_ERRORS = (TranslatorError, TokenizerError, IndexError, ValueError)


class FormulaShifter:
    """Um Translator do openpyxl reaproveitável: a fórmula é tokenizada uma vez e traduzida para
    quantos destinos for preciso. Erros (fórmula ilegível/intraduzível) viram ContractError."""

    def __init__(self, formula: str, origin: str, where: str = ""):
        self.origin, self._prefix = origin, (f"{where}!" if where else "")
        self._tr = self._err = None
        try:
            self._tr = Translator(formula, origin=origin)
        except _TRANSLATE_ERRORS as e:
            self._err = e                    # relatado em to(), com a célula de destino

    def to(self, dest: str) -> str:
        try:
            if self._err is not None:
                raise self._err
            return self._tr.translate_formula(dest)
        except _TRANSLATE_ERRORS as e:
            raise ContractError(f"{self._prefix}{dest}: não foi possível traduzir a fórmula "
                                f"de {self.origin} ({e})") from e


def translate_formula(formula: str, origin: str, dest: str, where: str = "") -> str:
    """Translator do openpyxl com os erros (fórmula intraduzível/ilegível) virando ContractError."""
    return FormulaShifter(formula, origin, where).to(dest)


def _check_xml_text(text: str, ref: str) -> None:
    """Fórmulas e erros não são limpos (mudaria o significado): caractere inválido -> ContractError."""
    if _ILLEGAL.search(text):
        raise ContractError(f"{ref}: texto com caracteres inválidos para o XML")
    try:
        etree.Element("x").text = text        # o lxml tem a palavra final sobre o que aceita
    except ValueError as e:
        raise ContractError(f"{ref}: texto com caracteres inválidos para o XML") from e


class DateStyles:
    """Diz se um índice de estilo (`s=`) é um formato de data, lendo styles.xml."""
    BUILTIN = set(range(14, 23)) | set(range(27, 37)) | {45, 46, 47} | set(range(50, 59))

    def __init__(self, styles_xml: bytes | None):
        self._flags: list[bool] = []
        if not styles_xml:
            return
        try:
            root = etree.fromstring(styles_xml)
        except etree.XMLSyntaxError as e:
            raise ContractError(f"xl/styles.xml ilegível ({e})") from e
        codes = {}
        numfmts = root.find(_q("numFmts"))
        if numfmts is not None:
            for e in numfmts.iterchildren(_q("numFmt")):
                codes[self._fmt_id(e.get("numFmtId"), "<numFmt>")] = e.get("formatCode", "")
        xfs = root.find(_q("cellXfs"))
        if xfs is not None:
            for i, xf in enumerate(xfs.iterchildren(_q("xf"))):
                fid = self._fmt_id(xf.get("numFmtId", "0"), f"<xf> nº {i} de cellXfs")
                self._flags.append(self._is_date_format(fid, codes.get(fid)))

    @staticmethod
    def _fmt_id(raw: str | None, where: str) -> int:
        if raw is None or not _ROWNUM.match(raw):
            raise ContractError(f"xl/styles.xml: {where} com numFmtId ausente ou inválido ({raw!r})")
        return int(raw)

    @classmethod
    def _is_date_format(cls, fid: int, code: str | None) -> bool:
        if fid in cls.BUILTIN:
            return True
        if code is None:
            return False
        stripped = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", code)
        return bool(re.search(r"[dmyhs]", stripped, re.I))

    def is_date(self, style) -> bool:
        i = int(style) if style not in (None, "") else 0
        return i < len(self._flags) and self._flags[i]


class SheetEditor:
    def __init__(self, xml: bytes, name: str = "aba"):
        self.name = name
        parser = etree.XMLParser(huge_tree=True, remove_blank_text=False)
        try:
            self.root = etree.fromstring(xml, parser)
        except etree.XMLSyntaxError as e:
            raise ContractError(f"{name}: XML da aba ilegível ({e})") from e
        self.sheet_data = self.root.find(_q("sheetData"))
        if self.sheet_data is None:
            raise ContractError(f"{name}: a aba não tem sheetData")
        self._rows = {}
        for row in self.sheet_data.iterchildren(_q("row")):
            r = self._row_number(row)
            if r in self._rows:
                raise ContractError(f"{name}: linha {r} repetida")
            self._rows[r] = row
            for c in row.iterchildren(_q("c")):
                m = _REF.match(c.get("r") or "")
                if m is None:
                    raise ContractError(f"{name}: célula sem referência válida (atributo r) na linha {r}")
                try:
                    column_index_from_string(m.group(1))
                except ValueError:
                    raise ContractError(f"{name}: célula {c.get('r')} com coluna além de XFD") from None
                if int(m.group(2)) != r:
                    raise ContractError(f"{name}: célula {c.get('r')} dentro da linha {r}")
        self._rownums = sorted(self._rows)
        # índice de células por linha, montado sob demanda: linha -> (colunas em ordem, coluna -> <c>)
        self._cols: dict[int, tuple[list[int], dict[int, etree._Element]]] = {}
        self._masters: dict[str, tuple[str, str]] | None = None
        # matrizes de várias células por linha coberta: linha -> [(ref, 1ª coluna, última coluna)]
        self._arrays: dict[int, list[tuple[str, int, int]]] | None = None
        self._shared: dict[str, FormulaShifter] = {}      # si -> tradutor do mestre (resolução das filhas)
        # (linha, coluna) -> (fórmula, tradutor ou None se ilegível, {linha first -> assinatura}):
        # cada fórmula é tokenizada uma vez para formula_family/is_foreign_formula; write invalida
        self._sigs: dict[tuple[int, int], tuple[str | None, FormulaShifter | None, dict]] = {}
        self.date_warnings: dict[str, int] = {}
        self.touched_rows: set[int] = set()
        self.autofilter_ref: str | None = None            # novo ref do <autoFilter>, se foi estendido

    def _row_number(self, row) -> int:
        r = row.get("r")
        if r is None or not _ROWNUM.match(r) or int(r) < 1:
            raise ContractError(f"{self.name}: linha sem o atributo r (ou com r inválido)")
        return int(r)

    # ---- navegação -------------------------------------------------------
    def row_exists(self, r: int) -> bool:
        return r in self._rows

    def _row(self, r: int, create: bool = True):
        row = self._rows.get(r)
        if row is None and create:
            row = etree.Element(_q("row"))
            row.set("r", str(r))
            pos = bisect.bisect_left(self._rownums, r)
            self.sheet_data.insert(pos, row)
            self._rownums.insert(pos, r)
            self._rows[r] = row
        return row

    def _col_index(self, r: int, row) -> tuple[list[int], dict]:
        idx = self._cols.get(r)
        if idx is None:
            cols, cells = [], {}
            for c in row.iterchildren(_q("c")):
                cc = column_index_from_string(_REF.match(c.get("r")).group(1))
                cols.append(cc)
                cells.setdefault(cc, c)
            idx = self._cols[r] = (cols, cells)
        return idx

    def _cell(self, r: int, col: int, create: bool = True):
        row = self._row(r, create)
        if row is None:
            return None
        cols, cells = self._col_index(r, row)
        c = cells.get(col)
        if c is not None or not create:
            return c
        # posição = nº de células com coluna menor (mesma regra da busca linear: as do modelo vêm em ordem)
        pos = next((i for i, cc in enumerate(cols) if cc > col), len(cols))
        cell = etree.Element(_q("c"))
        cell.set("r", f"{get_column_letter(col)}{r}")
        row.insert(pos, cell)
        cols.insert(pos, col)
        cells[col] = cell
        return cell

    def style_of(self, r: int, col: int) -> str | None:
        c = self._cell(r, col, create=False)
        return None if c is None else c.get("s")

    def nonempty_cells(self, rows, cols: set[int]) -> list[str]:
        found = []
        for r in rows:
            row = self._rows.get(r)
            if row is None:
                continue
            for c in row.findall(_q("c")):
                ci = column_index_from_string(_REF.match(c.get("r")).group(1))
                if ci in cols and any(c.find(_q(t)) is not None for t in ("v", "f", "is")):
                    found.append(c.get("r"))
        return found

    # ---- gravação ----------------------------------------------------------
    def write(self, r: int, col: int, value, *, style_from: tuple[int, int] | None = None,
              date_styles: DateStyles | None = None, replace_formula: bool = False) -> None:
        """Grava `value` em (r, col). None (ou texto que fica vazio depois de tirar os caracteres
        inválidos para o XML) não grava nada. Célula com fórmula do modelo é recusada, salvo com
        `replace_formula=True` (usado só pela extensão das fórmulas para trocar uma fórmula de outra
        família); mesmo assim o mestre de uma fórmula compartilhada nunca é trocado (as filhas
        perderiam a fórmula). Célula coberta por uma fórmula de matriz de VÁRIAS células
        (<f t="array" ref=...>) é sempre recusada: as filhas da matriz só têm valor e parecem
        constantes. Uma matriz de uma célula só é tratada como qualquer outra fórmula da célula.
        Todas as recusas são validadas antes de qualquer mudança (ver `check_write`).
        style_from: célula de onde copiar o estilo se a de destino não tiver um.
        date_styles: se dado, datas gravadas em célula sem formato de data são contadas por coluna.
        Para trocar mestres compartilhados junto com o grupo inteiro, ver `write_batch`."""
        self._write(r, col, value, style_from, date_styles, replace_formula, allow_master=False)

    def _write(self, r, col, value, style_from, date_styles, replace_formula, allow_master) -> None:
        content = self._validate(r, col, value, replace_formula, allow_master)
        if content is None:
            return
        c = self._cell(r, col)
        self._inherit_style(c, style_from)
        self._count_unformatted_date(c, col, value, date_styles)
        self._set(c, content)
        self._sigs.pop((r, col), None)
        self._rows[r].attrib.pop("spans", None)
        self.touched_rows.add(r)

    def check_write(self, r: int, col: int, value, *, replace_formula: bool = False) -> None:
        """Levanta o mesmo ContractError que `write(r, col, value, ...)` levantaria, sem mudar nada
        na planilha. Permite validar um lote inteiro antes da primeira gravação."""
        self._validate(r, col, value, replace_formula, allow_master=False)

    def write_batch(self, items, *, style_from: tuple[int, int] | None = None) -> None:
        """Grava o lote `items` = [(linha, coluna, valor), ...] trocando fórmulas existentes (como
        `write(..., replace_formula=True)`), de forma atômica: tudo é validado (`check_batch`) antes da
        primeira gravação. Diferente de `write`, pode trocar o MESTRE de uma fórmula compartilhada
        (<f t="shared" ref=... si=...>texto</f>), mas só quando TODAS as células do `ref` do grupo e
        todas as filhas com o mesmo si estão no lote: nenhuma filha fica sem mestre (o Excel repararia
        ou corromperia o arquivo). Usado pela fórmula completa da `linha_mestre` da produtora (build)."""
        items = list(items)
        self.check_batch(items)
        dropped = {si for _, _, si in self._batch_masters(items)}
        for r, col, value in items:
            self._write(r, col, value, style_from, None, True, allow_master=True)
        if dropped:                                       # grupos trocados por inteiro: sem mestre agora
            self._masters = None
            for si in dropped:
                self._shared.pop(si, None)

    def check_batch(self, items) -> None:
        """Levanta o ContractError que `write_batch(items)` levantaria, sem mudar nada na planilha."""
        items = list(items)
        for r, col, value in items:
            self._validate(r, col, value, True, allow_master=True)
        masters = self._batch_masters(items)
        if not masters:
            return
        planned = {(r, col) for r, col, _ in items}
        children: dict[str, list[tuple[int, int]]] = {}
        for f in self.sheet_data.iter(_q("f")):
            if f.get("t") == "shared" and not f.text:
                m = _REF.match(f.getparent().get("r"))
                children.setdefault(f.get("si"), []).append(
                    (int(m.group(2)), column_index_from_string(m.group(1))))
        for anchor, ref, si in masters:
            try:
                c1, r1, c2, r2 = range_boundaries(ref)
            except (ValueError, TypeError) as e:
                raise ContractError(f"{self.name}!{anchor}: fórmula compartilhada si={si} com ref "
                                    f"inválido {ref!r}") from e
            if None in (c1, r1, c2, r2):
                raise ContractError(f"{self.name}!{anchor}: fórmula compartilhada si={si} com ref "
                                    f"inválido {ref!r}")
            cells = [(rr, cc) for rr in range(r1, r2 + 1) for cc in range(c1, c2 + 1)]
            in_ref = set(cells)
            cells += [rc for rc in sorted(children.get(si, ())) if rc not in in_ref]
            left = [f"{get_column_letter(cc)}{rr}" for rr, cc in cells if (rr, cc) not in planned]
            if left:
                more = f" e mais {len(left) - 10}" if len(left) > 10 else ""
                raise ContractError(
                    f"{self.name}!{anchor} é o mestre da fórmula compartilhada si={si} (ref {ref}); "
                    "trocá-lo exige trocar todas as células do grupo no mesmo lote, mas ficariam de "
                    f"fora: {', '.join(left[:10])}{more}")

    def _batch_masters(self, items) -> list[tuple[str, str, str]]:
        """Mestres compartilhados entre as células do lote: [(célula, ref, si)]."""
        found = []
        for r, col, _ in items:
            c = self._cell(r, col, create=False)
            f = None if c is None else c.find(_q("f"))
            if f is not None and f.get("t") == "shared" and f.text:
                found.append((c.get("r"), f.get("ref") or c.get("r"), f.get("si")))
        return found

    def _validate(self, r: int, col: int, value, replace_formula: bool,
                  allow_master: bool = False) -> tuple[str | None, str, str] | None:
        """Recusas de `write` e conteúdo preparado; None = nada a gravar. Não muda a planilha.
        allow_master: só `write_batch`/`check_batch`, que conferem o grupo inteiro à parte."""
        if value is None or (isinstance(value, str) and not _ILLEGAL.sub("", value)):
            return None
        where = f"{self.name}!{get_column_letter(col)}{r}"
        array = self.array_formula_at(r, col)
        if array is not None:
            raise ContractError(f"{where} está dentro da fórmula de matriz {array}; não será sobrescrita")
        existing = self._cell(r, col, create=False)
        f = None if existing is None else existing.find(_q("f"))
        if f is not None:
            if not replace_formula:
                raise ContractError(f"{where} contém fórmula do modelo; não será sobrescrita")
            if f.get("t") == "shared" and f.text and not allow_master:
                raise ContractError(f"{where} é o mestre de uma fórmula compartilhada (si={f.get('si')}) "
                                    "de outra família; não será sobrescrito")
        return self._prepare(value, where)

    def _inherit_style(self, c, style_from: tuple[int, int] | None) -> None:
        if c.get("s") is not None or style_from is None:
            return
        ref = self._cell(style_from[0], style_from[1], create=False)
        if ref is not None and ref.get("s") is not None:
            c.set("s", ref.get("s"))

    def _count_unformatted_date(self, c, col: int, value, date_styles: DateStyles | None) -> None:
        if (date_styles is not None and isinstance(value, (dt.datetime, dt.date))
                and not date_styles.is_date(c.get("s"))):
            letter = get_column_letter(col)
            self.date_warnings[letter] = self.date_warnings.get(letter, 0) + 1

    @staticmethod
    def _prepare(value, ref: str) -> tuple[str | None, str, str]:
        """Valida e converte o valor sem tocar na planilha: (atributo t, tag, texto).
        `ref` é a célula nas mensagens de erro (aba!célula)."""
        if isinstance(value, Formula):
            text = str(value)
            kind, tag, text = None, "f", (text[1:] if text.startswith("=") else text)
        elif isinstance(value, XlError):
            kind, tag, text = "e", "v", str(value)
        elif isinstance(value, bool):
            kind, tag, text = "b", "v", "1" if value else "0"
        elif isinstance(value, (dt.datetime, dt.date)):
            try:
                kind, tag, text = None, "v", _num(to_serial(value), ref)
            except ContractError as e:
                raise ContractError(f"{ref}: {e}") from None
        elif isinstance(value, (int, float)):
            kind, tag, text = None, "v", _num(value, ref)
        elif isinstance(value, str):
            kind, tag, text = "inlineStr", "t", _ILLEGAL.sub("", value)
        else:
            raise ContractError(f"tipo de valor não suportado em {ref}: {type(value).__name__}")
        _check_xml_text(text, ref)
        return kind, tag, text

    @staticmethod
    def _set(c, content: tuple[str | None, str, str]) -> None:
        """Troca o conteúdo da célula pelo já validado em `_prepare`."""
        kind, tag, text = content
        for child in list(c):
            c.remove(child)
        c.attrib.pop("t", None)
        if kind is not None:
            c.set("t", kind)
        if tag == "t":
            el = etree.SubElement(etree.SubElement(c, _q("is")), _q("t"))
            el.text = text
            el.set(f"{{{XML_NS}}}space", "preserve")
        else:
            etree.SubElement(c, _q(tag)).text = text

    # ---- fórmulas ------------------------------------------------------------
    def _build_masters(self) -> None:
        self._masters = {}
        for c in self.sheet_data.iter(_q("c")):
            f = c.find(_q("f"))
            if f is not None and f.get("t") == "shared" and f.text:
                self._masters[f.get("si")] = (c.get("r"), f.text)

    def _array_index(self) -> dict[int, list[tuple[str, int, int]]]:
        """Índice (montado uma vez) das fórmulas de matriz de VÁRIAS células: linha coberta ->
        [(ref, 1ª coluna, última coluna)]. Matriz de uma célula só (ref = a própria célula) fica de
        fora: o <f> está na própria célula, que é tratada como qualquer fórmula. ref inválido em
        qualquer matriz -> ContractError."""
        if self._arrays is None:
            self._arrays = {}
            for f in self.sheet_data.iter(_q("f")):
                if f.get("t") != "array":
                    continue
                anchor = f.getparent().get("r")
                ref = f.get("ref") or anchor
                try:
                    bounds = range_boundaries(ref)
                except (ValueError, TypeError) as e:
                    raise ContractError(f"{self.name}: fórmula de matriz com ref inválido {ref!r}") from e
                if None in bounds:
                    raise ContractError(f"{self.name}: fórmula de matriz com ref inválido {ref!r}")
                c1, r1, c2, r2 = bounds
                if (c1, r1, c2, r2) == range_boundaries(anchor):
                    continue                      # uma célula só: a própria âncora
                for rr in range(r1, r2 + 1):
                    self._arrays.setdefault(rr, []).append((ref, c1, c2))
        return self._arrays

    def array_formula_at(self, r: int, col: int) -> str | None:
        """ref da fórmula de matriz de várias células (<f t="array" ref=...>) que cobre (r, col), ou
        None (inclusive para a âncora de uma matriz de uma célula só, ver `_array_index`)."""
        for ref, c1, c2 in self._array_index().get(r, ()):
            if c1 <= col <= c2:
                return ref
        return None

    def formula_at(self, r: int, col: int) -> str | None:
        """Texto da fórmula da célula (com '='), resolvendo fórmulas compartilhadas.
        Filha compartilhada sem mestre ou fórmula intraduzível -> ContractError com a célula."""
        c = self._cell(r, col, create=False)
        if c is None:
            return None
        f = c.find(_q("f"))
        if f is None:
            return None
        if f.get("t") == "shared" and not f.text:
            if self._masters is None:
                self._build_masters()
            si = f.get("si")
            shifter = self._shared.get(si)
            if shifter is None:
                master = self._masters.get(si)
                if master is None:
                    raise ContractError(f"{self.name}!{c.get('r')}: fórmula compartilhada "
                                        f"si={si} sem célula mestre")
                origin, text = master
                shifter = self._shared[si] = FormulaShifter("=" + text, origin, where=self.name)
            return shifter.to(c.get("r"))
        return "=" + (f.text or "")

    def text_at(self, r: int, col: int) -> str | None:
        """Texto de uma célula `inlineStr` (o tipo que este módulo grava); None nos demais casos.
        Usado na cobertura para conferir o rótulo `TOTAL` da BASE."""
        c = self._cell(r, col, create=False)
        if c is None or c.get("t") != "inlineStr":
            return None
        is_ = c.find(_q("is"))
        return None if is_ is None else "".join(is_.itertext())

    def _at_first(self, r: int, col: int, first: int, failed: list[str] | None = None) -> str | None:
        """Fórmula de (r, col) traduzida para a linha `first` (a "assinatura" da família), ou None se
        a célula não tem fórmula, tem <f/> vazio, é ilegível (não se traduz nem para a própria célula;
        filha compartilhada sem mestre) ou sai da planilha em `first` (ex.: =COUNTA(AP3:AP11) na 12
        com first=3). As ilegíveis são anotadas em `failed`. O resultado fica em cache por célula
        (ver `_formula_entry`)."""
        f, shifter, sigs = self._formula_entry(r, col)
        if f is None:
            return None
        letter = get_column_letter(col)
        if shifter is None:
            if failed is not None:
                failed.append(f"{letter}{r}")
            return None
        if first not in sigs:
            try:
                sigs[first] = shifter.to(f"{letter}{first}")
            except ContractError:
                sigs[first] = None
        return sigs[first]

    def _formula_entry(self, r: int, col: int) -> tuple[str | None, FormulaShifter | None, dict]:
        """(fórmula, tradutor, cache de assinaturas) da célula, tokenizando a fórmula uma vez só.
        fórmula None = sem fórmula ou <f/> vazio; tradutor None = ilegível (não se traduz nem para a
        própria célula, ou filha compartilhada sem mestre). Invalidado por `write` na célula."""
        entry = self._sigs.get((r, col))
        if entry is None:
            origin = f"{get_column_letter(col)}{r}"
            try:
                f = self.formula_at(r, col)
            except ContractError:                  # filha compartilhada sem mestre/intraduzível
                f, shifter = "", None
            else:
                xf = None if not f else self._cell(r, col, create=False).find(_q("f"))
                if not f or f == "=":
                    f = shifter = None
                elif xf.get("t") == "shared" and not xf.text:
                    # filha legível: traduzir o texto dela ou traduzir o mestre para o mesmo destino
                    # dá o mesmo resultado (a tradução só soma o deslocamento a cada referência), e
                    # o tradutor do mestre já está tokenizado (ver formula_at)
                    shifter = self._shared[xf.get("si")]
                else:
                    try:
                        shifter = FormulaShifter(f, origin, where=self.name)
                        shifter.to(origin)
                    except ContractError:
                        shifter = None
            entry = self._sigs[(r, col)] = (f, shifter, {})
        return entry

    def _row_invariant(self, r: int, col: int) -> bool:
        """A fórmula de (r, col) não muda ao descer uma linha (só referências absolutas ou nenhuma)."""
        f, shifter, _ = self._formula_entry(r, col)
        if shifter is None:
            return False
        try:
            return shifter.to(f"{get_column_letter(col)}{r + 1}") == f
        except ContractError:
            return False

    def formula_family(self, col: int, first: int = 1) -> tuple[int, str] | None:
        """Família das fórmulas de dados da coluna: (última linha >= `first` da família, assinatura).

        Assinatura = a fórmula traduzida para a linha `first`. A referência é a primeira fórmula de
        dados (a mais alta >= `first` que se traduz para `first`); outra célula é da MESMA FAMÍLIA se
        a sua fórmula, traduzida para `first`, dá a mesma assinatura. Linhas de resumo/rodapé (ex.:
        =COUNTA(AP3:AP39) na 40) são de outra família e nunca viram base de extensão (para decidir
        se uma fórmula DENTRO das linhas carregadas deve ser trocada, ver `is_foreign_formula`).
        None = não há fórmula de dados na coluna. Se não há fórmula de dados mas há fórmula ilegível
        (não se traduz nem para a própria célula, ou filha compartilhada sem mestre), levanta
        ContractError nomeando essas células. As assinaturas ficam em cache por célula (cada fórmula
        é tokenizada uma vez para esta função e `is_foreign_formula`).

        Limites conhecidos (o modelo de setembro não tem nenhum dos dois casos):
        - sem nenhuma fórmula de dados, um resumo traduzível (só referências absolutas) seria tomado
          como dado;
        - uma família cujas fórmulas apontam mais linhas acima do que há entre `first` e a 1ª fórmula
          (ex.: =A{r-3} começando na linha 6 com first=3) não tem referência: nenhuma se traduz para
          `first` (cairia em A0) e o resultado é None, o que bloqueia a carga ("classificação
          ausente") em vez de estender errado. Uma vez achada a referência, todo membro da família se
          traduz de volta para `first` (a tradução só desfaz o deslocamento), então a comparação
          pelas assinaturas não perde membros."""
        rows = [r for r in self._rownums if r >= first and self.has_formula_element(r, col)]
        failed: list[str] = []
        ref = next(((r, t) for r in rows if (t := self._at_first(r, col, first, failed)) is not None), None)
        if ref is None:
            if failed:                             # há fórmulas, mas só ilegíveis ou de resumo
                letter = get_column_letter(col)
                raise ContractError(
                    f"{self.name}!{letter}: fórmula(s) ilegível(is)/intraduzível(is) em {letter} "
                    f"({', '.join(failed)}); nenhuma fórmula de dados para estender")
            return None
        last = next(r for r in reversed(rows) if r == ref[0] or self._at_first(r, col, first) == ref[1])
        return last, ref[1]

    def is_foreign_formula(self, r: int, col: int, first: int, signature: str) -> bool:
        """(r, col) tem um <f> que NÃO é fórmula de dados e deve ser trocado pela da família
        `signature` (ver `formula_family`)? Sim quando o <f> é vazio, ilegível (inclusive filha
        compartilhada sem mestre), sai da planilha ao ser traduzido para `first` (resumo relativo,
        ex.: =COUNTA(AP3:AP11) na linha 12 com first=3) ou tem assinatura diferente E não muda de
        uma linha para outra (resumo só com referências absolutas, ex.: =$A$1&"x" ou
        =SUBTOTAL(3,$AP$3:$AP$11)): uma fórmula assim não pode ser fórmula de dados por linha.

        Fórmula de assinatura diferente COM referências relativas fica intacta: é tratada como uma
        VARIANTE de dados do modelo (nas colunas sem `linha_mestre` da produtora). Em
        DESPESAS!AP o modelo de setembro traz duas versões (AP881:AP1070, mestre compartilhado si=22,
        tem um ramo a mais "Transferência entre Empresas Coligadas" -> "Desconsiderar"); por decisão do
        usuário de 13/10/2026 a coluna é configurada e a versão completa da linha 881 é aplicada a
        todas as linhas (ver build._unify_formulas): lá esta função só separa, na contagem, as fórmulas
        de outra família (overwritten_constants) das versões de dados trocadas (unified_formula_rows).

        Limite conhecido (falso negativo): um resumo com referências relativas cuja tradução para
        `first` ainda cabe na planilha passa por variante e NÃO é substituído nem contado. Um resumo
        na linha R cuja referência relativa mais alta é a linha s cabe quando s - (R - first) >= 1.
        Ex.: =SUM(AP$3:AP11) na 12 (AP$3 é absoluta na linha); =COUNTA(B25:B39) na 40 com
        first >= 16 (com first=3 vira B-12:B2 e é reconhecida); =COUNTA(B38:B39) na 40 com first=3
        (vira B1:B2). Nos `first` reais (2, 3 e 5) um resumo sobre a faixa inteira de dados (de
        `first` até a linha anterior, s = first) sai da planilha sempre que a faixa tem pelo menos
        `first` linhas (R >= 2*first), o que vale para as bases reais (centenas de linhas): é
        reconhecido."""
        if not self.has_formula_element(r, col):
            return False
        sig = self._at_first(r, col, first)
        if sig is None:
            return True
        return sig != signature and self._row_invariant(r, col)

    def formula_rows(self, col: int, first: int = 1) -> list[int]:
        """Linhas >= `first`, em ordem, cuja célula da coluna tem um elemento <f> (qualquer tipo)."""
        return [r for r in self._rownums if r >= first and self.has_formula_element(r, col)]

    def signature_at(self, r: int, col: int, first: int) -> str | None:
        """Fórmula de (r, col) traduzida para a linha `first` (None: sem fórmula, <f/> vazio, ilegível
        ou fora da planilha em `first`). Em cache por célula; ver `formula_family`."""
        return self._at_first(r, col, first)

    def shared_group(self, r: int, col: int) -> tuple[str, bool] | None:
        """(si, é o mestre?) se (r, col) tem fórmula compartilhada; None nos demais casos."""
        c = self._cell(r, col, create=False)
        f = None if c is None else c.find(_q("f"))
        if f is None or f.get("t") != "shared":
            return None
        return f.get("si"), bool(f.text)

    def last_formula_row(self, col: int, first: int = 1) -> int | None:
        """Última linha >= `first` com fórmula da família de dados (ver `formula_family`)."""
        family = self.formula_family(col, first)
        return None if family is None else family[0]

    def has_formula_element(self, r: int, col: int) -> bool:
        """A célula tem um elemento <f> (qualquer tipo, mesmo vazio)?"""
        c = self._cell(r, col, create=False)
        return c is not None and c.find(_q("f")) is not None

    # ---- filtro ------------------------------------------------------------
    def extend_autofilter(self, last_row: int) -> str | None:
        """Estende o <autoFilter> da aba até `last_row` se o filtro termina antes; nunca encolhe.
        Devolve o novo ref (ou None se não há filtro ou nada mudou); fica em `autofilter_ref`."""
        af = self.root.find(_q("autoFilter"))
        if af is None or not af.get("ref"):
            return None
        try:
            c1, r1, c2, r2 = range_boundaries(af.get("ref"))
        except (ValueError, TypeError) as e:
            raise ContractError(f"{self.name}: autoFilter com ref inválido {af.get('ref')!r}") from e
        if last_row <= r2:
            return None
        self.autofilter_ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{last_row}"
        af.set("ref", self.autofilter_ref)
        return self.autofilter_ref

    # ---- serialização ------------------------------------------------------
    def to_bytes(self) -> bytes:
        self._update_dimension()
        return etree.tostring(self.root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _update_dimension(self) -> None:
        dim = self.root.find(_q("dimension"))
        if dim is None:
            return
        lo_r = lo_c = 10 ** 9
        hi_r = hi_c = 0
        for row in self.sheet_data.findall(_q("row")):
            for c in row.findall(_q("c")):
                m = _REF.match(c.get("r"))
                ci, ri = column_index_from_string(m.group(1)), int(m.group(2))
                lo_r, hi_r = min(lo_r, ri), max(hi_r, ri)
                lo_c, hi_c = min(lo_c, ci), max(hi_c, ci)
        if not hi_r:
            return
        try:
            c1, r1, c2, r2 = range_boundaries(dim.get("ref") or "")
        except (ValueError, TypeError) as e:
            raise ContractError(f"{self.name}: dimension com ref inválido {dim.get('ref')!r}") from e
        if None in (c1, r1, c2, r2):
            raise ContractError(f"{self.name}: dimension com ref inválido {dim.get('ref')!r}")
        dim.set("ref", f"{get_column_letter(min(c1, lo_c))}{min(r1, lo_r)}:"
                       f"{get_column_letter(max(c2, hi_c))}{max(r2, hi_r)}")


def _read_xml(zf: zipfile.ZipFile, name: str):
    try:
        return etree.fromstring(zf.read(name))
    except KeyError:
        raise ContractError(f"o pacote não tem {name}") from None
    except etree.XMLSyntaxError as e:
        raise ContractError(f"{name} ilegível ({e})") from e


def _sheet_parts(zf: zipfile.ZipFile) -> dict[str, str]:
    """Nome da aba (NFC) -> parte do zip, na ordem do workbook.xml."""
    wb = _read_xml(zf, "xl/workbook.xml")
    rels = _read_xml(zf, "xl/_rels/workbook.xml.rels")
    targets = {r.get("Id"): r.get("Target") for r in rels}
    sheets = wb.find(_q("sheets"))
    if sheets is None:
        raise ContractError("xl/workbook.xml sem <sheets>")
    parts = {}
    for sheet in sheets:
        name = unicodedata.normalize("NFC", sheet.get("name") or "")
        target = targets.get(sheet.get(f"{{{R_NS}}}id"))
        if not target:
            raise ContractError(f"a aba {name!r} não tem parte no pacote (relação sem alvo)")
        parts[name] = target.lstrip("/") if target.startswith("/") else "xl/" + target
    return parts


class WorkbookSession:
    """Abre um .xlsx modelo, edita só as abas pedidas e grava uma cópia. Use com `with`."""

    def __init__(self, template):
        self.template = Path(template)
        try:
            self._zin = zipfile.ZipFile(self.template)
        except (OSError, zipfile.BadZipFile) as e:
            raise ContractError(f"não foi possível abrir o modelo {self.template.name}: {e}") from e
        try:
            self.parts = _sheet_parts(self._zin)
            self._editors: dict[str, SheetEditor] = {}
            names = self._zin.namelist()
            self.date_styles = DateStyles(
                self._zin.read("xl/styles.xml") if "xl/styles.xml" in names else None)
        except BaseException:
            self._zin.close()
            raise

    def editor(self, sheet_name: str) -> SheetEditor:
        key = unicodedata.normalize("NFC", sheet_name)
        if key not in self.parts:
            raise ContractError(f"o modelo não tem a aba {sheet_name!r}")
        if key not in self._editors:
            try:
                xml = self._zin.read(self.parts[key])
            except KeyError:
                raise ContractError(f"a aba {sheet_name!r} aponta para {self.parts[key]}, "
                                    "que não existe no pacote") from None
            self._editors[key] = SheetEditor(xml, key)
        return self._editors[key]

    def save(self, out_path) -> Path:
        """Grava a cópia editada. Tudo é montado e validado antes de tocar no destino; o zip é escrito
        num temporário da mesma pasta e só substitui `out_path` no sucesso (uma saída anterior
        sobrevive a qualquer falha). Gravar por cima do próprio modelo é recusado."""
        out_path = Path(out_path)
        if out_path.resolve() == self.template.resolve():
            raise ContractError("a saída não pode ser o próprio modelo")
        replaced = {self.parts[k]: ed.to_bytes() for k, ed in self._editors.items()}
        replaced.update(self._plumbing())
        out_path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(prefix=f".{out_path.name}.", suffix=".tmp", dir=out_path.parent)
        os.close(fd)
        try:
            with zipfile.ZipFile(tmp, "w") as zout:
                for info in self._zin.infolist():
                    if info.filename == "xl/calcChain.xml":
                        continue
                    data = replaced.get(info.filename)
                    if data is None:
                        data = self._zin.read(info.filename)
                    zi = zipfile.ZipInfo(info.filename, info.date_time)
                    zi.compress_type = info.compress_type
                    zi.external_attr = info.external_attr
                    zout.writestr(zi, data)
            os.replace(tmp, out_path)
        except BaseException:
            Path(tmp).unlink(missing_ok=True)
            raise
        return out_path

    def _plumbing(self) -> dict[str, bytes]:
        """Partes de ligação do pacote que mudam: workbook.xml (recálculo ao abrir) e, se houver
        calcChain, as referências a ele em [Content_Types].xml e workbook.xml.rels."""
        names = set(self._zin.namelist())
        order = list(self.parts)                      # localSheetId = posição da aba em <sheets>
        filters = {order.index(k): ed.autofilter_ref for k, ed in self._editors.items() if ed.autofilter_ref}
        workbook = _sync_filter_names(self._zin.read("xl/workbook.xml"), filters)
        out = {"xl/workbook.xml": _force_full_calc(workbook)}
        if "xl/calcChain.xml" in names:
            for name, pattern in (
                    ("[Content_Types].xml", rb'<Override[^>]*PartName="/xl/calcChain\.xml"[^>]*/>'),
                    ("xl/_rels/workbook.xml.rels", rb'<Relationship[^>]*Target="[^"]*calcChain\.xml"[^>]*/>')):
                if name in names:
                    out[name] = re.sub(pattern, b"", self._zin.read(name))
        return out

    def close(self) -> None:
        self._zin.close()

    def __enter__(self) -> "WorkbookSession":
        return self

    def __exit__(self, *exc) -> None:
        self.close()


_CALCPR = re.compile(rb"<(?:[A-Za-z_][\w.-]*:)?calcPr\b[^>]*>")
_FULLCALC = re.compile(rb"""\sfullCalcOnLoad\s*=\s*(?:"[^"]*"|'[^']*')""")


def _force_full_calc(data: bytes) -> bytes:
    """Garante `fullCalcOnLoad="1"` no <calcPr> (qualquer ordem de atributos, aspas simples ou
    duplas, com ou sem prefixo de namespace; idempotente)."""
    m = _CALCPR.search(data)
    if m is None:
        raise ContractError("workbook.xml sem <calcPr>; não é possível forçar o recálculo ao abrir")
    tag = m.group(0)
    if _FULLCALC.search(tag):
        new = _FULLCALC.sub(b' fullCalcOnLoad="1"', tag, count=1)
    else:
        new = re.sub(rb"^(<(?:[A-Za-z_][\w.-]*:)?calcPr)", rb'\1 fullCalcOnLoad="1"', tag, count=1)
    return data[:m.start()] + new + data[m.end():]


_DEFINED_NAME = re.compile(rb"<((?:[A-Za-z_][\w.-]*:)?)definedName\b([^>]*)>([^<]*)</\1definedName>")
_FILTER_NAME = re.compile(rb"""\bname\s*=\s*["']_xlnm\._FilterDatabase["']""")
_LOCAL_ID = re.compile(rb"""\blocalSheetId\s*=\s*["'](\d+)["']""")


def _sync_filter_names(data: bytes, filters: dict[int, str]) -> bytes:
    """Acompanha o <autoFilter> estendido no nome oculto _xlnm._FilterDatabase da mesma aba
    (localSheetId), trocando só o intervalo depois do '!' e sem nunca encolhê-lo. Aceita o
    <definedName> com ou sem prefixo de namespace (ex.: <x:definedName>). Ficam byte a byte: outros
    nomes, nomes com texto que não seja um intervalo simples (ex.: #REF!) e nomes com várias áreas
    (texto com ',', ex.: DESPESAS!$A$2:$B$9,DESPESAS!$D$2:$D$9 — alargar uma área só seria arbitrário;
    por simplicidade isso vale também para uma aba cujo nome tenha vírgula).
    Se o nome oculto já diverge do <autoFilter> (outras colunas/linhas), o resultado é a UNIÃO dos
    dois intervalos, de propósito: nunca encolhe o que o modelo já filtrava e passa a cobrir as
    linhas carregadas."""
    if not filters:
        return data

    def fix(m: re.Match) -> bytes:
        ns, attrs, text = m.group(1), m.group(2), m.group(3)
        if b"," in text:
            return m.group(0)
        lid = _LOCAL_ID.search(attrs)
        if not _FILTER_NAME.search(attrs) or lid is None or int(lid.group(1)) not in filters:
            return m.group(0)
        prefix, sep, rng = text.rpartition(b"!")
        try:
            old = range_boundaries(rng.decode("ascii").replace("$", ""))
        except (ValueError, TypeError, UnicodeDecodeError):
            return m.group(0)
        if not sep or None in old:
            return m.group(0)
        new = range_boundaries(filters[int(lid.group(1))])
        c1, r1 = min(old[0], new[0]), min(old[1], new[1])
        c2, r2 = max(old[2], new[2]), max(old[3], new[3])
        if (c1, r1, c2, r2) == old:
            return m.group(0)
        ref = f"${get_column_letter(c1)}${r1}:${get_column_letter(c2)}${r2}".encode("ascii")
        return b"<" + ns + b"definedName" + attrs + b">" + prefix + sep + ref + b"</" + ns + b"definedName>"

    return _DEFINED_NAME.sub(fix, data)


def changed_parts(template, output) -> dict[str, list[str]]:
    """Compara dois .xlsx parte a parte (bytes)."""
    with zipfile.ZipFile(template) as a, zipfile.ZipFile(output) as b:
        an, bn = set(a.namelist()), set(b.namelist())
        changed = sorted(n for n in an & bn if a.read(n) != b.read(n))
    return {"changed": changed, "removed": sorted(an - bn), "added": sorted(bn - an)}
