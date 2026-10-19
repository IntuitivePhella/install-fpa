from __future__ import annotations

from pathlib import Path

import openpyxl

from ..config import DEST, PESQUISA, R_BLANK, R_FOOTER, R_NUCLEO, R_PESQ, RRA_FOOTER
from ..headers import check_header, load_headers, normalize
from ..produtoras import Regras, regras_padrao
from ..readers import convert_xls_to_xlsx
from ..types import Removal, Sanitized
from ._xlsx import RowCopier, blank, check_header_extent

REPORT = "rra"
NCOLS = 49
HEADER_ROW = 3
FIRST_RAW_ROW = 4


def sanitize_rra(xls_path, *, work_dir, expected_header: list[str] | None = None,
                 regras: Regras | None = None) -> Sanitized:
    """XLS do Esferas -> XLSX (Excel COM, preserva fórmulas) -> sanitização."""
    xlsx = convert_xls_to_xlsx(xls_path, Path(work_dir) / (Path(xls_path).stem + ".xlsx"))
    return sanitize_rra_xlsx(xlsx, expected_header=expected_header, regras=regras)


def sanitize_rra_xlsx(xlsx_path, *, expected_header: list[str] | None = None,
                      regras: Regras | None = None) -> Sanitized:
    regras = regras or regras_padrao()
    wf = openpyxl.load_workbook(xlsx_path)
    wv = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wf["Plan1"] if "Plan1" in wf.sheetnames else wf.worksheets[0]
    wsv = wv[ws.title]
    header_cells = ws[HEADER_ROW]
    header = [normalize(c.value) if c.value is not None else "" for c in header_cells[:NCOLS]]
    check_header(REPORT, header, expected_header or load_headers()[REPORT])
    check_header_extent(REPORT, header_cells, NCOLS)
    i_nucleo = header.index("Núcleo")
    i_titulo = header.index("Título do job")
    first_dest = DEST[REPORT]["first_row"]

    out = Sanitized(report=REPORT, header=header, rows=[], db_rows=[], source_rows=[],
                    raw_count=0, title=ws["A1"].value)
    copier = RowCopier(REPORT)

    # iter_rows com max_col fixo: `ws[r]` recalcula max_column a cada linha (quadrático).
    last = ws.max_row
    pairs = zip(ws.iter_rows(min_row=FIRST_RAW_ROW, max_row=last, max_col=NCOLS),
                wsv.iter_rows(min_row=FIRST_RAW_ROW, max_row=last, max_col=NCOLS))
    for r, (cells, vcells) in enumerate(pairs, start=FIRST_RAW_ROW):
        values = [c.value for c in cells]
        out.raw_count += 1
        if blank(values):
            out.technical.append(Removal(r, R_BLANK))
            continue
        if str(values[0]).strip().upper() == RRA_FOOTER:     # rodapé: técnico, valores não lidos
            out.technical.append(Removal(r, R_FOOTER))
            continue
        nucleo = normalize(values[i_nucleo]) if values[i_nucleo] is not None else ""
        if nucleo != regras.nucleo_rra:
            out.exclusions.append(Removal(r, R_NUCLEO))
            continue
        if PESQUISA in str(values[i_titulo] or "").upper():
            out.exclusions.append(Removal(r, R_PESQ))
            continue

        dest = first_dest + len(out.rows)
        row, db_row = copier.copy(cells, vcells, r, dest, values[0])
        out.rows.append(row)
        out.db_rows.append(db_row)
        out.source_rows.append(r)
    out.warnings.extend(copier.warnings())
    return out
